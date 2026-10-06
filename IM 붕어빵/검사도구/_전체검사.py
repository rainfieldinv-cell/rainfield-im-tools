# -*- coding: utf-8 -*-
"""가진 IM 원본을 **전부** 돌려서 한눈에 비교한다.

★한 문서만 보고 기준값을 정하면 다른 IM 에서 반드시 어긋난다. 고칠 때마다 이걸 돌려서
  어느 문서도 나빠지지 않았는지 확인할 것.

사용: python _전체검사.py
"""
import os
import sys

# ★이 파일은 '검사도구' 폴더 안에 있다. 프로그램(pdf_to_ppt)과 자료(_원본.pdf)는
#   옆 폴더에 있으므로 찾아갈 길을 열어 준다.
_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, _ROOT)   # 2026-10-06 한 폴더로 합침: 엔진(pdf_to_ppt)이 바로 윗 폴더에 있다
_자료 = os.path.join(_ROOT, "자료")
import glob
import os
import tempfile

import fitz
from pptx import Presentation

import pdf_to_ppt as P

HERE = _자료
EMU = 914400 / 72.0

# ★검사용 원본은 IM 폴더의 **IM(샘플)** 한 곳에 모아 둔다(2026-10-06, 예전 자료\검사용원본\). 바탕화면이나 다운로드에 둔 것을
#   가리키면 그게 지워지는 순간 검사가 통째로 멈춘다(실제로 한 번 겪었다).
#   새 IM 을 받으면 그 폴더에 복사만 하면 자동으로 검사 대상에 들어온다.
SOURCES = sorted(glob.glob(os.path.join(os.path.dirname(_ROOT), "IM(샘플)", "*.pdf")))


def one(path):
    raw = open(path, "rb").read()
    out = os.path.join(tempfile.gettempdir(), "_all.pptx")
    info = P.convert(raw, out)
    doc = fitz.open(stream=raw, filetype="pdf")

    n_tab = sum(s["표"] for s in info["pages"])
    n_box = sum(s["글상자"] for s in info["pages"])
    n_rect = sum(s["네모"] for s in info["pages"])

    # ★안쪽 표를 몇 개나 **따로 떼어냈나** — 이건 렌더 그림으로는 절대 안 보인다.
    #   겉모습은 붙어 있으나 떨어져 있으나 똑같고, 파워포인트에서 클릭해야 안다.
    #   그래서 반드시 숫자로 세야 한다(7쪽 계좌개설 표가 이렇게 붙은 채 지나갔다).
    inner = 0
    for i in range(doc.page_count):
        _r, plan, _b = P._plan(doc[i], P._lines_of(doc[i]))
        inner += sum(len(t["skip"]) for t in plan)

    # 표 위에 뜬 글상자(=칸에 못 들어간 글자)
    prs = Presentation(out)
    over = 0
    for sl in prs.slides:
        tabs = [(sh.left / EMU, sh.top / EMU, (sh.left + sh.width) / EMU,
                 (sh.top + sh.height) / EMU)
                for sh in sl.shapes if getattr(sh, "has_table", False) and sh.has_table]
        for sh in sl.shapes:
            if sh.shape_type == 17 and sh.has_text_frame and sh.text_frame.text.strip():
                t = sh.text_frame.text.strip()
                if P._UNIT_RE.match(t):
                    continue        # 단위는 일부러 따로 뺀 것이다 — 문제가 아니다
                x = (sh.left + sh.width / 2) / EMU
                y = (sh.top + sh.height / 2) / EMU
                if any(a <= x <= c and b <= y <= d for a, b, c, d in tabs):
                    over += 1

    import _선검사
    import _정렬검사
    miss = sum(_선검사.check(i + 1, doc[i]) for i in range(doc.page_count))
    bad = cells = 0
    for i in range(doc.page_count):
        a, b = _정렬검사.check(i + 1, doc[i])
        bad += a
        cells += b
    doc.close()
    return {"쪽": info["count"], "표": n_tab, "글상자": n_box, "네모": n_rect,
            "안쪽표": inner,
            "표위글상자": over, "못넣은선": miss, "어긋난칸": bad, "칸수": cells}


if __name__ == "__main__":
    import contextlib
    import io as _io
    rows = []
    for p in SOURCES:
        if not os.path.exists(p):
            continue
        buf = _io.StringIO()
        try:
            with contextlib.redirect_stdout(buf):
                r = one(p)
        except Exception as e:
            print(f"{os.path.basename(p)[:30]:<32} 실패 {type(e).__name__}: {e}")
            continue
        rows.append((os.path.basename(p), r))
    print(f"{'원본':<32}{'쪽':>4}{'표':>5}{'안쪽표':>7}{'네모':>6}{'글상자':>7}"
          f"{'표위글상자':>9}{'못넣은선':>8}{'어긋난칸':>10}")
    for name, r in rows:
        print(f"{name[:30]:<32}{r['쪽']:>4}{r['표']:>5}{r['안쪽표']:>7}{r['네모']:>6}"
              f"{r['글상자']:>7}{r['표위글상자']:>9}{r['못넣은선']:>8}"
              f"{r['어긋난칸']:>6}/{r['칸수']:<5}")
