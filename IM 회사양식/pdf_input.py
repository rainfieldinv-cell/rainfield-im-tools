# -*- coding: utf-8 -*-
"""PDF 원본 → 회사 양식 규격에 맞춘 워드 본문.

2026-10-08 사용자 확정 규칙
  · 결과는 **회사 양식 규격(A4·여백) 고정**, 머리말·꼬리말 위치·크기도 양식 고정.
  · **원본 1쪽 = 결과 1쪽.** 원본 쪽 내용이 양식 안 칸보다 크면 **비율 그대로 줄여** 넣는다(작으면 그대로).
  · 1쪽(증권사 표지)은 빼고, 마지막 쪽이 증권사 연락처면 뺀다.
  · 쪽마다 같은 자리에 되풀이되는 위·아래 띠(증권사 머리말·로고·쪽번호)는 지운다.

순서
  1. 뺄 쪽·띠 지우기(fitz)  2. 쪽마다 남은 내용의 범위를 잰다
  3. pdf2docx 로 워드(원본 한 쪽 = 한 구역)  4. 구역마다 양식 규격으로 바꾸고 내용 전체를 같은 비율로 줄인다
  5. pdf2docx 가 실패해 건너뛴 쪽은 **원본 쪽을 고화질 그림으로** 넣는다(내용이 빠지면 안 된다)
"""
import io
import logging
import os
import re
import tempfile
from collections import Counter

import fitz

EMAIL_RE = re.compile(r"[\w.\-]+@[\w\-]+\.[\w.]+")
PHONE_RE = re.compile(r"0\d{1,2}[-)\s.]\d{3,4}[-\s.]\d{4}")

TOP_BAND = 0.16       # 쪽 높이의 위 16% 를 머리말 띠로 본다(12% 로는 머리말 밑줄이 빠져 빈 표로 남았다 — 등촌역)
BOTTOM_BAND = 0.12    # 아래 12% 를 꼬리말 띠로 본다
REPEAT = 0.5          # 절반 넘는 쪽에 같은 자리로 나오면 '되풀이' 로 본다

# 회사 양식(형식_1) 규격 — cm
PAGE_W, PAGE_H = 21.0, 29.7
TOP_CM, BOTTOM_CM, LEFT_CM, RIGHT_CM = 3.0, 2.5, 2.5, 2.5
HEADER_CM, FOOTER_CM = 1.7, 1.2
BOX_W_PT = (PAGE_W - LEFT_CM - RIGHT_CM) / 2.54 * 72        # 본문 칸 가로(pt)
BOX_H_PT = (PAGE_H - TOP_CM - BOTTOM_CM) / 2.54 * 72        # 본문 칸 세로(pt)
HEIGHT_SAFETY = 0.92  # 글꼴이 Pretendard 로 바뀌면 줄이 조금 높아진다 → 세로는 8% 여유(5% 로는 등촌역 쪽 몇 개가 넘쳤다)

PAGE_NO_RE = re.compile(r"[-–\s]*(?:page\s*)?\d{1,3}(?:\s*(?:/|of)\s*\d{1,3})?[-–\s]*", re.I)
HIGHLIGHT_RE = re.compile(r"executive\s*summary|investment\s*highlight|투자\s*하이라이트|투자\s*포인트|"
                          r"핵심\s*요약|하이라이트|highlights?", re.I)


# ──────────────────────────────────────────────
# 원본 살펴보기
# ──────────────────────────────────────────────
def text_per_page(pdf_bytes):
    """쪽당 글자 수(글자를 그림으로 바꾼 PDF 를 미리 알아보려고)."""
    doc = fitz.open(stream=pdf_bytes, filetype="pdf")
    n = sum(len(p.get_text().strip()) for p in doc)
    return n / max(doc.page_count, 1)


def last_page_is_contact(doc):
    t = doc[-1].get_text()
    return len(set(EMAIL_RE.findall(t))) + len(set(PHONE_RE.findall(t))) >= 2


def find_highlight_pages(pdf_bytes):
    """하이라이트(Executive Summary 등) 쪽 번호(1부터) 목록. 표지 다음 6쪽 안에서, 쪽 위쪽 글로 찾는다."""
    doc = fitz.open(stream=pdf_bytes, filetype="pdf")
    hits = []
    for i in range(1, min(doc.page_count, 7)):
        top = doc[i].get_text("text", clip=fitz.Rect(0, 0, doc[i].rect.width, doc[i].rect.height * 0.35))
        if HIGHLIGHT_RE.search(top):
            hits.append(i + 1)
        elif hits:
            break
    return hits


BOND_NEED_PT = 270    # 사모사채 개요(제목 + 9줄 표)에 드는 높이 어림(약 9.5cm)


def bond_place(rep, highlight_pages):
    """사모사채 개요 자리 — 하이라이트 마지막 쪽 뒤(2026-10-08 사용자 확정).

    그 쪽에 개요가 들어갈 여백이 남으면 (k, True) = 이어서, 아니면 (k, False) = 새 쪽.
    하이라이트를 못 찾았으면 None = 표지 바로 다음 쪽.
    """
    if not highlight_pages:
        return None
    last = max(highlight_pages)
    for k, p in enumerate(rep.get("쪽", [])):
        if p["원본쪽"] == last:
            return (k, BOX_H_PT - p["높이"] >= BOND_NEED_PT)
    return None


def page_png(pdf_bytes, pno, zoom=0.6):
    doc = fitz.open(stream=pdf_bytes, filetype="pdf")
    return doc[pno].get_pixmap(matrix=fitz.Matrix(zoom, zoom)).tobytes("png")


def suggest_title(pdf_bytes):
    """1쪽(표지)에서 가장 큰 글씨 줄들 → 표지 제목 후보(최대 2줄)."""
    doc = fitz.open(stream=pdf_bytes, filetype="pdf")
    lines = []
    for b in doc[0].get_text("dict")["blocks"]:
        for l in b.get("lines", []):
            t = re.sub(r"\s+", " ", "".join(s["text"] for s in l["spans"])).strip()
            if not t or len(t) > 60 or re.search(r"memorandum|confidential|private|securities|division", t, re.I):
                continue
            if re.fullmatch(r"[\d.\-/\s]+", t):
                continue
            lines.append((max(s["size"] for s in l["spans"]), t))
    if not lines:
        return []
    top = max(s for s, _ in lines)
    out = []
    for s, t in lines:
        if s >= top * 0.8 and t not in out:
            out.append(t)
    return out[:2]


def all_photos(pdf_bytes):
    """원본에 든 그림 **전부**(로고처럼 여러 쪽에 되풀이되는 것·아주 작은 것만 뺀다).

    돌려주는 것: [{'blob','w','h','page','score'}] — 표지에 쓸 만한 순(색이 다양한 큰 실사 사진 먼저).
    blob 은 원본에 든 해상도 그대로(가능하면 원본 파일 그대로) 꺼낸다.
    """
    from PIL import Image
    doc = fitz.open(stream=pdf_bytes, filetype="pdf")
    pages_of = Counter()
    for page in doc:
        for x in {img[0] for img in page.get_images(full=True)}:
            pages_of[x] += 1
    seen, out = set(), []
    for pno, page in enumerate(doc):
        for img in page.get_images(full=True):
            xref, smask = img[0], img[1]
            if xref in seen or pages_of[xref] >= 3:
                continue
            seen.add(xref)
            try:
                if smask:
                    pix = fitz.Pixmap(doc, xref)
                    if pix.n - pix.alpha >= 4:
                        pix = fitz.Pixmap(fitz.csRGB, pix)
                    pix = fitz.Pixmap(pix, fitz.Pixmap(doc, smask))
                    blob = pix.tobytes("png")
                else:
                    ext = doc.extract_image(xref)
                    blob = ext["image"]
                    if ext.get("ext") not in ("jpeg", "jpg", "png"):
                        pix = fitz.Pixmap(doc, xref)
                        if pix.n - pix.alpha >= 4:
                            pix = fitz.Pixmap(fitz.csRGB, pix)
                        blob = pix.tobytes("png")
                im = Image.open(io.BytesIO(blob))
                if im.mode == "CMYK":
                    im = im.convert("RGB")
                    b = io.BytesIO()
                    im.save(b, "JPEG", quality=95)
                    blob = b.getvalue()
                w, h = im.size
                if w < 120 or h < 80:
                    continue
                rgb = im.convert("RGB")
                colors = len(set(rgb.resize((80, 80)).getdata()))
                bright = sum(rgb.convert("L").resize((60, 60)).getdata()) / 3600
            except Exception:
                continue
            ratio = w / h
            score = (colors > 1500, bright < 210, 1.2 <= ratio <= 2.6, w * h)
            out.append({"blob": blob, "w": w, "h": h, "page": pno + 1, "score": score})
    out.sort(key=lambda d: d["score"], reverse=True)
    return out


# ──────────────────────────────────────────────
# 띠 지우기 · 내용 범위 재기
# ──────────────────────────────────────────────
EDGE_TOP, EDGE_BOTTOM = 0.055, 0.05   # 쪽 맨 가장자리 여백 — 본문이 안 들어오는 자리


def _is_page_no(txt, pno):
    """쪽번호인가. 숫자만 덜렁 있으면 그 쪽의 실제 번호와 거의 같을 때만(표 숫자 '266' 을 지운 일이 있다)."""
    if not PAGE_NO_RE.fullmatch(txt):
        return False
    if re.search(r"/|of|page|페이지|^[-–]|[-–]$", txt.strip(), re.I):
        return True
    nums = re.findall(r"\d+", txt)
    return len(nums) == 1 and pno - 3 <= int(nums[0]) <= pno + 3


def _band_items(page):
    """위·아래 띠 안의 글줄·그림·선을 (종류, 반올림한 자리, 내용, 범위) 로."""
    h = page.rect.height
    pno = page.number + 1
    top, bot = h * TOP_BAND, h * (1 - BOTTOM_BAND)
    out = []
    for b in page.get_text("dict")["blocks"]:
        # ★글은 '줄' 단위로 본다 — 덩어리(block) 단위면 머리말이 아래 본문과 한 덩어리로 묶여 띠 밖으로
        #   잡혀 하나도 안 지워졌다(고산3지구 머리말 249줄 그대로)
        if b["type"] == 0:
            units = [(fitz.Rect(l["bbox"]), "".join(s["text"] for s in l["spans"]).strip()) for l in b["lines"]]
        else:
            # ★그림은 '내용' 까지 같아야 되풀이로 본다 — 성수동은 제목 숫자('2.', '8.')가 그림 조각인데,
            #   자리·크기만 보니 숫자가 달라도 같은 그림으로 쳐서 제목 숫자가 지워졌다
            import hashlib
            digest = hashlib.md5(b.get("image", b"")[:4096]).hexdigest()[:10]
            units = [(fitz.Rect(b["bbox"]), "\x00" + digest)]
        for r, txt in units:
            if not (r.y1 <= top or r.y0 >= bot):
                continue
            if txt is not None and txt.startswith("\x00"):
                out.append(("i", (round(r.x0 / 6), round(r.y0 / 6), round(r.width / 6)), txt, r))
                continue
            if txt is not None:
                if not txt:
                    continue
                # 가로 쪽은 높이가 낮아 5% 가 33pt 밖에 안 된다(성수동 'Ⅰ. Project 개요' 가 남았다) → 최소 40pt
                # 아래쪽 가장자리는 안 쓴다 — 쪽 바닥까지 내려온 표 숫자('33,389')를 지운 일이 있다(등촌역).
                #   아래 꼬리말은 '되풀이' 와 쪽번호 규칙으로 잡힌다.
                edge = r.y1 <= max(h * EDGE_TOP, 40)
                if edge or _is_page_no(txt.replace("페이지", "page"), pno):
                    # 쪽번호('13 / 14', '- 3 -', 'Page 3')는 자릿수에 따라 자리가 조금씩 달라 '되풀이' 로
                    # 안 잡혔다(등촌역 13·14쪽에 남음) → 띠 안에 있으면 무조건 지운다.
                    # 쪽 맨 가장자리 여백의 글(고산 2쪽 오른쪽 위 머리말)도 한 번만 나와도 지운다.
                    out.append(("pn", (), "", r))
                    continue
                key = re.sub(r"\d+", "#", txt)            # 쪽번호는 숫자만 다르다
                if re.fullmatch(r"\s*\d{1,3}\s*", txt):
                    # 숫자만 있는 줄 — 쪽번호를 본문부터 1로 다시 세는 문서가 있어(고산3지구 5쪽의 '1')
                    # '높이' 가 여러 쪽에서 같으면 쪽번호로 본다(아래 strip_repeating_bands 에서 센다)
                    out.append(("num", (round(r.y0 / 4),), "", r))
                    continue
                if len(key.strip()) < 4:
                    continue                               # 'Ⅲ.' '개요' 같은 짧은 제목 조각은 머리말로 안 본다(성수동)
                out.append(("t", (round(r.x0 / 6), round(r.y0 / 6)), key, r))
            else:
                out.append(("i", (round(r.x0 / 6), round(r.y0 / 6), round(r.width / 6)), "", r))
    for d in page.get_drawings():
        r = d["rect"]
        if r.y1 <= top or r.y0 >= bot:
            out.append(("d", (round(r.x0 / 6), round(r.y0 / 6), round(r.width / 6)), "", r))
    return out


def _safe_text_targets(doc, pno, rects):
    """글 줄 지우기를 사본 쪽에서 먼저 해 보고, 다른 본문 글까지 사라지게 하는 줄은 빼고 돌려준다.

    ★PDF 지우기는 글자 '모양 상자' 가 겹치면 지운다. 글꼴 정보가 엉터리로 큰 PDF 에서는 꼬리말 한 줄을
      지우는데 그 위 표 숫자까지 사라졌다(등촌역 8쪽: '오이코스자산운용 대체투자본부' 를 지우니 '33,389' 도).
    반환: (지워도 되는 범위들, PDF 에선 못 지운 줄들의 글 — 워드로 바꾼 뒤 지운다)
    """
    page = doc[pno]
    targets = {tuple(r) for r in rects}
    keep_txt = Counter(t for r, t in _lines(page) if tuple(r) not in targets)
    ok, skipped = [], []
    for r in rects:
        tmp = fitz.open()
        tmp.insert_pdf(doc, from_page=pno, to_page=pno)
        tp = tmp[0]
        tp.add_redact_annot(r, fill=False)
        _apply(tp, images=False, graphics=None, text=True)
        after = Counter(t for rr, t in _lines(tp))
        lost = keep_txt - after
        if lost:
            skipped.append(next((t for rr, t in _lines(page) if tuple(rr) == tuple(r)), ""))
        else:
            ok.append(r)
    return ok, [t for t in skipped if t]


def strip_repeating_bands(doc, skipped=None):
    """여러 쪽에 같은 자리로 되풀이되는 위·아래 띠 요소를 지운다. 지운 쪽 수를 돌려준다.

    skipped(list) 를 주면 PDF 에서 안전하게 못 지운 머리말·꼬리말 글을 거기 모은다(워드로 바꾼 뒤 지운다).
    """
    per_page = [_band_items(p) for p in doc]
    cnt = Counter()
    for items in per_page:
        cnt.update({(k, pos, key) for k, pos, key, _ in items})
    need = max(2, int(doc.page_count * REPEAT))
    rep = {sig for sig, c in cnt.items() if c >= need}
    # ★같은 머리말 글이라도 쪽마다 자리가 조금 다를 수 있다(고산3지구 2쪽 DISCLAIMER 의 오른쪽 위 머리말).
    #   → 글은 '자리' 를 빼고 '내용' 만으로도 되풀이를 센다.
    txt_cnt = Counter()
    for items in per_page:
        txt_cnt.update({key for k, pos, key, _ in items if k == "t" and len(key) >= 8})
    rep_txt = {key for key, c in txt_cnt.items() if c >= need}
    num_cnt = Counter()
    for items in per_page:
        num_cnt.update({pos for k, pos, key, _ in items if k == "num"})
    touched = 0
    for page, items in zip(doc, per_page):
        hit = [(k, r) for k, pos, key, r in items
               if k == "pn" or (k, pos, key) in rep or (k == "t" and key in rep_txt)
               or (k == "num" and num_cnt[pos] >= max(3, need // 2))]
        if not hit:
            continue
        # ★두 번에 나눠 지운다. 한 번에 지우면 선·그림 자리에 걸친 **글자까지** 지워진다
        #   (성수동: 쪽마다 같은 자리의 제목 밑줄을 지우다 그 위 제목 '2. 개발계획' 이 같이 사라졌다).
        #   ① 글 줄 → 글만  ② 선·그림 → 그 도형만(글은 그대로)
        # fill=False — 흰 네모로 덮으면 pdf2docx 가 그 네모를 빈 표로 만든다(실제: 등촌역 쪽마다 빈 표)
        texts = [r for k, r in hit if k in ("t", "pn", "num")]
        # ★선·그림은 그 쪽 본문 글보다 위(위 띠) / 아래(아래 띠)에 있을 때만 머리말·꼬리말로 본다.
        #   제목 아래 밑줄처럼 본문 사이에 끼어 있으면 쪽마다 같은 자리여도 본문 장식이다(성수동 제목 밑줄).
        hit_txt = set(map(tuple, texts))
        body = [r for r, t in _lines(page) if tuple(r) not in hit_txt]
        h = page.rect.height
        top_body = min((r.y0 for r in body if r.y0 < h * 0.5), default=h)
        bot_body = max((r.y1 for r in body if r.y1 > h * 0.5), default=0)
        shapes = [r for k, r in hit if k not in ("t", "pn", "num")
                  and ((r.y1 <= h * 0.5 and r.y1 <= top_body + 2) or (r.y0 > h * 0.5 and r.y0 >= bot_body - 2))]
        if texts:
            texts, miss = _safe_text_targets(doc, page.number, texts)
            if skipped is not None:
                skipped.extend(miss)
            for r in texts:
                page.add_redact_annot(r, fill=False)
            if texts:
                _apply(page, images=False, graphics=None, text=True)
        # ★선은 '걸치기만 해도' 지운다 — 원본 머리말 밑줄이 쪽 바깥(x=-1)까지 그어져 있어 '완전히 든 것만' 으로는
        #   안 지워졌다(등촌역). 본문보다 위·아래에 있는 것만 골랐으므로 본문 선은 건드리지 않는다.
        #   그림은 하나씩 따로 지운다 — 딱 붙은 두 조각을 한 번에 지우니 하나가 남았다(등촌역 머리말 장식).
        for r in shapes:
            page.add_redact_annot(r + (-1, -1, 1, 1), fill=False)
            _apply(page, images=True, graphics="touched", text=False)
        touched += 1
    return touched


ROLE_RE = re.compile(r"담당|팀장|부장|차장|과장|대리|주임|사원|이사|상무|전무|본부|부서|팀\b|연락처|문의|"
                     r"tel|e-?mail|mobile|fax|contact|휴대|전화|이메일", re.I)


def _lines(page):
    out = []
    for b in page.get_text("dict")["blocks"]:
        for l in b.get("lines", []):
            t = "".join(s["text"] for s in l["spans"]).strip()
            if t:
                out.append((fitz.Rect(l["bbox"]), t))
    return out


def strip_contacts(page):
    """원본 증권사 담당자 연락처 묶음(이름·직급·전화·메일)을 지운다 — 어느 쪽에 있든(2026-10-08 사용자 확정).

    메일·전화번호가 2개 이상인 줄들을 씨앗으로, 위아래로 바짝 붙은(30pt 안) 이름·직급·'담당 부서' 줄까지
    넓힌 범위의 글·선·배경·그림을 지운다. 떨어져 있는 본문 문단은 건드리지 않는다.
    (실제: 고산3지구 2쪽 DISCLAIMER 아래 '금융조달 담당 부서 – 키움증권 …' 4명 연락처)
    반환: 지웠으면 True
    """
    lines = _lines(page)
    hits = [(r, t) for r, t in lines if EMAIL_RE.search(t) or PHONE_RE.search(t)]
    found = set()
    for _, t in hits:
        found |= set(EMAIL_RE.findall(t)) | set(PHONE_RE.findall(t))
    if len(found) < 2:
        return False
    box = fitz.Rect()
    for r, _ in hits:
        box |= r
    grew = True
    while grew:
        grew = False
        for r, t in lines:
            if r in box or box.contains(r):
                continue
            near = (0 <= box.y0 - r.y1 <= 30) or (0 <= r.y0 - box.y1 <= 30) or r.intersects(box)
            if near and (ROLE_RE.search(t) or len(t) <= 14):
                box |= r
                grew = True
    box = box + (-8, -6, 8, 6)
    # 연락처 묶음에 걸친 얇은 띠·선(제목 줄 회색 배경, 칸 선)도 함께 — 큰 테두리는 건드리지 않는다
    for d in page.get_drawings():
        r = d["rect"]
        if r.intersects(box) and r.height < 40 and r.width < page.rect.width * 0.95:
            box |= r
    page.add_redact_annot(box, fill=False)
    _apply(page, images=True, graphics="covered", text=True)
    return True


def strip_band_logos(page):
    """위·아래 머리말 띠 안의 작은 그림(로고)을 지운다 — 한 쪽에만 있어도(되풀이 여부 상관없이)."""
    h, w = page.rect.height, page.rect.width
    n = 0
    for info in page.get_image_info():
        r = fitz.Rect(info["bbox"])
        in_band = r.y1 <= h * TOP_BAND or r.y0 >= h * (1 - BOTTOM_BAND)
        if in_band and r.width < w * 0.35 and r.height < h * 0.10:
            page.add_redact_annot(r + (-1, -1, 1, 1), fill=False)
            n += 1
    if n:
        _apply(page, images=True, graphics=None, text=False)
    return n


def _apply(page, images=True, graphics="covered", text=True):
    """지우기 실행. graphics: None(선 안 건드림) / 'covered'(범위에 완전히 든 선만) / 'touched'(걸친 선도)."""
    kw = {"images": fitz.PDF_REDACT_IMAGE_REMOVE if images else fitz.PDF_REDACT_IMAGE_NONE,
          "graphics": {"covered": fitz.PDF_REDACT_LINE_ART_REMOVE_IF_COVERED,
                       "touched": fitz.PDF_REDACT_LINE_ART_REMOVE_IF_TOUCHED}.get(graphics, fitz.PDF_REDACT_LINE_ART_NONE),
          "text": fitz.PDF_REDACT_TEXT_REMOVE if text else fitz.PDF_REDACT_TEXT_NONE}
    try:
        page.apply_redactions(**kw)
    except TypeError:                                       # 옛 판에는 graphics·text 인자가 없다
        page.apply_redactions(images=kw["images"])


def content_box(page):
    """띠를 지운 뒤 쪽에 남은 내용(글·그림·선)의 범위."""
    r = fitz.Rect()
    for b in page.get_text("dict")["blocks"]:
        if b["type"] == 0 and not "".join(s["text"] for l in b["lines"] for s in l["spans"]).strip():
            continue
        r |= fitz.Rect(b["bbox"])
    for d in page.get_drawings():
        dr = d["rect"]
        if dr.width > 1 or dr.height > 1:
            r |= dr
    for info in page.get_image_info():
        r |= fitz.Rect(info["bbox"])
    return r & page.rect if not r.is_empty else fitz.Rect(0, 0, 1, 1)


# ──────────────────────────────────────────────
# 워드로 바꾸고 규격 맞추기
# ──────────────────────────────────────────────
class _CatchIgnored(logging.Handler):
    def __init__(self):
        super().__init__()
        self.pages = []

    def emit(self, record):
        m = re.search(r"Ignore page (\d+)", record.getMessage())
        if m:
            self.pages.append(int(m.group(1)))


def convert(pdf_bytes, drop_cover=True, drop_last=None):
    """PDF bytes → (docx bytes, 보고 dict).

    보고['쪽'] = 결과 구역 순서대로 원본 쪽 번호(1부터)·줄인 비율·쓰인 높이(pt)·그림으로 넣었는지.
    """
    import docx
    doc = fitz.open(stream=pdf_bytes, filetype="pdf")
    n = doc.page_count
    rep = {"원본 쪽수": n}
    if drop_last is None:
        drop_last = n > 2 and last_page_is_contact(doc)
    keep = list(range(n))
    if drop_last and n > 2:
        keep.remove(n - 1)
    if drop_cover and n > 1:
        keep.remove(0)
    rep["원본 끝쪽 뺌"], rep["원본 표지 뺌"] = bool(drop_last and n > 2), bool(drop_cover and n > 1)
    leftover = []
    rep["머리말·꼬리말 지운 쪽"] = strip_repeating_bands(doc, leftover)
    # 원본 담당자 연락처(어느 쪽이든)·머리말 띠의 로고(한 쪽에만 있어도) — 2026-10-08 사용자 확정
    rep["연락처 지운 쪽"] = [i + 1 for i in keep if strip_contacts(doc[i])]
    rep["로고 지운 쪽"] = [i + 1 for i in keep if strip_band_logos(doc[i])]
    boxes = [content_box(doc[i]) for i in range(n)]
    sizes = [(doc[i].rect.width, doc[i].rect.height) for i in range(n)]
    doc.select(keep)

    from pdf2docx import Converter
    import docx.table as _dt
    catcher = _CatchIgnored()
    logging.getLogger().addHandler(catcher)
    # ★pdf2docx 는 표 칸 합치기가 하나라도 겹치면 **그 쪽을 통째로 버린다**('Ignore page …').
    #   안산 8쪽('정보통신업 사업자 관련 설명' 표)·33쪽이 원본 그대로도 이렇게 빠졌다(2026-10-08).
    #   → 바꾸는 동안만, 합치기가 안 되는 칸 하나는 합치지 않고 넘어가게 해 표 나머지를 살린다.
    _orig_merge = _dt._Cell.merge

    def _safe_merge(self, other):
        try:
            return _orig_merge(self, other)
        except Exception:
            rep.setdefault("못 합친 칸", 0)
            rep["못 합친 칸"] += 1
            return self
    _dt._Cell.merge = _safe_merge
    try:
        with tempfile.TemporaryDirectory() as td:
            src, dst = os.path.join(td, "in.pdf"), os.path.join(td, "out.docx")
            doc.save(src)
            cv = Converter(src)
            try:
                cv.convert(dst)
            finally:
                cv.close()
            with open(dst, "rb") as f:
                raw = f.read()
    finally:
        logging.getLogger().removeHandler(catcher)
        _dt._Cell.merge = _orig_merge
    ignored = {keep[p - 1] for p in catcher.pages if 0 < p <= len(keep)}

    d = docx.Document(io.BytesIO(raw))
    rep["워드에서 지운 머리말 줄"] = drop_paragraphs(d, set(leftover))
    rep["줄바꿈 탭 고침"] = fix_wrap_tabs(d)
    rep["기호 글꼴 고침"] = fix_symbol_runs(d)
    done = [k for k in keep if k not in ignored]
    page_info = []
    for (secs, grps), k in zip(_pages(d), done):
        s, used = _fit_page(secs, grps, boxes[k])
        page_info.append({"원본쪽": k + 1, "비율": s, "높이": used, "그림": False})
    # 실패해 건너뛴 쪽 → 원본 쪽을 그림으로(제자리에)
    for k in sorted(ignored):
        idx = sum(1 for p in done if p < k) + sum(1 for p in ignored if p < k)
        img = fitz.open(stream=pdf_bytes, filetype="pdf")[k].get_pixmap(matrix=fitz.Matrix(3, 3), clip=boxes[k])
        used = _insert_picture_page(d, idx, img.tobytes("png"), boxes[k])
        page_info.insert(idx, {"원본쪽": k + 1, "비율": None, "높이": used, "그림": True})
    rep["쪽"] = page_info
    rep["그림으로 넣은 쪽"] = [k + 1 for k in sorted(ignored)]
    out = io.BytesIO()
    d.save(out)
    return out.getvalue(), rep


def _section_groups(d):
    """구역(=원본 한 쪽)마다 블록 묶음. 구역은 '문단 안 sectPr' 로 끝난다(마지막은 본문 sectPr)."""
    from docx.oxml.ns import qn
    groups, cur = [], []
    for el in d.element.body:
        if el.tag not in (qn("w:p"), qn("w:tbl")):
            continue
        cur.append(el)
        if el.tag == qn("w:p") and el.find(qn("w:pPr") + "/" + qn("w:sectPr")) is not None:
            groups.append(cur)
            cur = []
    groups.append(cur)
    return groups


def drop_paragraphs(d, texts):
    """PDF 에서 안전하게 못 지운 머리말·꼬리말 줄을 워드에서 지운다(그 글과 똑같은 문단만, 표 밖만)."""
    from docx.oxml.ns import qn
    if not texts:
        return 0
    want = {re.sub(r"\s+", "", t) for t in texts}
    n = 0
    # 머리말·꼬리말은 늘 그 쪽(구역)의 맨 처음·맨 끝에 있다 → 구역마다 글 있는 문단 앞뒤 3개만 본다
    #   (숫자만 있는 쪽번호 '30' 이 본문 어딘가의 '30' 문단까지 지우지 않게)
    for grp in _section_groups(d):
        paras = [e for e in grp if e.tag == qn("w:p") and "".join(x.text or "" for x in e.iter(qn("w:t"))).strip()]
        for p in paras[:3] + paras[-3:]:
            if p.getparent() is None:
                continue
            t = re.sub(r"\s+", "", "".join(x.text or "" for x in p.iter(qn("w:t"))))
            if t and t in want and p.find(qn("w:pPr") + "/" + qn("w:sectPr")) is None:
                p.getparent().remove(p)
                n += 1
    return n


def fix_wrap_tabs(d):
    """pdf2docx 가 '줄이 바뀌는 자리' 에 넣은 탭을 띄어쓰기로 바꾸고, 둘째 줄부터 들여쓰기로 대신한다.

    ★원본은 글머리 문단의 둘째 줄을 글머리 뒤에 맞추려고 [첫 줄 글][탭][둘째 줄 글] 로 만들어져 있다.
      글꼴이 Pretendard 로 바뀌어 첫 줄이 짧아지면 탭이 같은 줄 오른쪽으로 튀어
      '…법인 요건을          갖출 예정' 처럼 벌어졌다(안산 8쪽, 2026-10-08).
    조건: 문단에 직접 정한 탭 자리가 1~3개이고 모두 왼쪽 들여쓰기 근처(2cm 안)에 있으며, 탭 앞에 글이 있을 때.
      연달아 있는 탭(안산 '[부가통신사업의 정의]' 는 2개)은 띄어쓰기 하나로. 둘째 줄부터는 가장 먼 탭 자리에.
    """
    from docx.oxml.ns import qn
    n = 0
    for p in d.element.body.iter(qn("w:p")):
        ppr = p.find(qn("w:pPr"))
        tabs = ppr.find(qn("w:tabs")) if ppr is not None else None
        stops = tabs.findall(qn("w:tab")) if tabs is not None else []
        if not 1 <= len(stops) <= 3:
            continue
        ind = ppr.find(qn("w:ind"))
        left = int(ind.get(qn("w:left"), "0")) if ind is not None and (ind.get(qn("w:left")) or "0").lstrip("-").isdigit() else 0
        poss = [int(t.get(qn("w:pos"), "0")) for t in stops]
        pos = max(poss)
        if pos <= left or pos > left + 1134:
            continue
        seen_text, changed, prev_tab = False, False, False
        for r in p.iter(qn("w:r")):
            for c in list(r):
                if c.tag == qn("w:t") and (c.text or "").strip() and not (c.text or "").strip() in "▪•·∙-–※✓➢◦○●■□▶▷":
                    seen_text, prev_tab = True, False
                elif c.tag == qn("w:br"):
                    prev_tab = False
                elif c.tag == qn("w:tab") and seen_text:
                    if not prev_tab:                       # 연달아 있는 탭은 띄어쓰기 하나로
                        t = r.makeelement(qn("w:t"), {"{http://www.w3.org/XML/1998/namespace}space": "preserve"})
                        t.text = " "
                        c.addprevious(t)
                    r.remove(c)
                    changed, prev_tab = True, True
        if changed:
            if ind is None:
                ind = ppr.makeelement(qn("w:ind"), {})
                after = next((ppr.find(qn(t)) for t in ("w:jc", "w:rPr", "w:sectPr")
                              if ppr.find(qn(t)) is not None), None)       # ind 는 jc·rPr·sectPr 보다 앞
                if after is not None:
                    after.addprevious(ind)
                else:
                    ppr.append(ind)
            ind.attrib.pop(qn("w:firstLine"), None)
            ind.set(qn("w:left"), str(pos))
            ind.set(qn("w:hanging"), str(pos - left))
            n += 1
    return n


_SYMBOL_FONT_RE = re.compile(r"wingding|webding|symbol|marlett", re.I)


def fix_symbol_runs(d):
    """기호 글꼴(Wingdings 등)이 붙어 있는데 글자는 보통 유니코드 기호인 조각 → 'Segoe UI Symbol'.

    ★pdf2docx 가 '▪'(U+25AA)에 Wingdings 를 붙여 와서, 워드가 Wingdings 의 같은 번호 모양인 '✦' 로
      그렸다(안산). Wingdings 전용 글자(U+F000~F0FF)인 조각은 그대로 둔다.
    """
    from docx.oxml.ns import qn
    n = 0
    for r in d.element.body.iter(qn("w:r")):
        rpr = r.find(qn("w:rPr"))
        rf = rpr.find(qn("w:rFonts")) if rpr is not None else None
        if rf is None or not any(_SYMBOL_FONT_RE.search(rf.get(qn("w:" + k)) or "") for k in ("ascii", "hAnsi", "eastAsia", "cs")):
            continue
        txt = "".join(t.text or "" for t in r.findall(qn("w:t")))
        if txt and all(not (0xF000 <= ord(c) <= 0xF0FF) for c in txt):
            for k in ("ascii", "hAnsi", "eastAsia", "cs"):
                rf.set(qn("w:" + k), "Segoe UI Symbol")
            n += 1
    return n


def _pages(d):
    """구역을 '원본 한 쪽' 단위로 묶는다 → [(구역들, 블록 묶음들)].

    ★pdf2docx 는 한 쪽 안에서도 단(칼럼) 나누기에 구역을 여러 개 쓴다(새 쪽 → 2단 이어서 → 다음 단).
      구역 하나 = 한 쪽으로 보고 전부 '새 쪽' 으로 바꿨더니 빈 쪽이 생겼다(등촌역 7쪽이 구역 3개).
      → '새 쪽으로 시작하는 구역' 에서만 쪽을 나눈다.
    """
    from docx.enum.section import WD_SECTION
    secs, grps = list(d.sections), _section_groups(d)
    out = []
    for i, (sec, grp) in enumerate(zip(secs, grps)):
        if i == 0 or sec.start_type in (WD_SECTION.NEW_PAGE, WD_SECTION.ODD_PAGE, WD_SECTION.EVEN_PAGE):
            out.append(([], []))
        out[-1][0].append(sec)
        out[-1][1].append(grp)
    return out


def _set_geometry(sec, new_page=True):
    from docx.shared import Cm
    from docx.enum.section import WD_SECTION, WD_ORIENT
    sec.orientation = WD_ORIENT.PORTRAIT
    sec.page_width, sec.page_height = Cm(PAGE_W), Cm(PAGE_H)
    sec.top_margin, sec.bottom_margin = Cm(TOP_CM), Cm(BOTTOM_CM)
    sec.left_margin, sec.right_margin = Cm(LEFT_CM), Cm(RIGHT_CM)
    sec.header_distance, sec.footer_distance = Cm(HEADER_CM), Cm(FOOTER_CM)
    if new_page:
        sec.start_type = WD_SECTION.NEW_PAGE


def _fit_page(secs, grps, box):
    """원본 한 쪽(구역 1개 이상)을 양식 규격으로: 내용을 같은 비율로 줄이고(크면만) 위쪽 빈 자리를 없앤다.

    반환: (비율, 결과에서 쓰인 높이 pt)
    """
    from docx.oxml.ns import qn
    first = secs[0]
    old_left, old_top = first.left_margin.pt, first.top_margin.pt
    old_w = first.page_width.pt - first.left_margin.pt - first.right_margin.pt
    s = min(1.0, BOX_W_PT / max(old_w, 1), BOX_H_PT * HEIGHT_SAFETY / max(box.height, 1))
    for i, sec in enumerate(secs):
        _set_geometry(sec, new_page=(i == 0))
        cols = sec._sectPr.find(qn("w:cols"))               # 2단 같은 단 너비·간격도 같은 비율로
        if cols is not None and s < 0.999:
            for el in [cols] + cols.findall(qn("w:col")):
                for a in ("space", "w"):
                    v = el.get(qn("w:" + a))
                    if v and v.isdigit():
                        el.set(qn("w:" + a), str(int(int(v) * s)))
    for grp in grps:
        for el in grp:
            scale_xml(el, s, old_left)
    lead = max(0.0, (box.y0 - old_top) * s)                 # 위 띠를 지운 자리 = 쓸데없는 빈 칸
    trim_leading(grps[0], lead)
    for grp in grps[:-1]:                                    # 쪽 안의 구역 끝(단 나눔) 문단도 납작하게
        if grp:
            _flatten_holder(grp[-1]) if grp[-1].find(qn("w:pPr") + "/" + qn("w:sectPr")) is not None else None
    trim_trailing(grps[-1])
    return s, box.height * s


# ── 같은 비율로 줄이기 ──
_TWIPS = {"spacing": ("before", "after"), "ind": ("left", "right", "firstLine", "hanging", "start", "end"),
          "tab": ("pos",), "tblW": ("w",), "tcW": ("w",), "gridCol": ("w",), "tblInd": ("w",),
          "trHeight": ("val",), "tblCellSpacing": ("w",)}
_HALF = {"sz", "szCs", "kern", "position"}
_EMU = {"extent": ("cx", "cy"), "ext": ("cx", "cy"), "off": ("x", "y"),
        "effectExtent": ("l", "t", "r", "b")}


def scale_xml(root, s, old_left=0.0):
    """root 안의 길이(글자 크기·간격·들여쓰기·표 너비·그림 크기·위치)를 모두 s 배로."""
    if s >= 0.999:
        return
    W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    WP = "{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}"
    A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"

    def mul(el, attr, ns=W, minimum=None):
        k = ns + attr if ns else attr
        v = el.get(k)
        if v is None or not v.lstrip("-").isdigit():
            return
        nv = int(round(int(v) * s))
        if minimum is not None:
            nv = max(minimum, nv)
        el.set(k, str(nv))

    for el in root.iter():
        tag = el.tag
        if not isinstance(tag, str):
            continue
        if tag.startswith(W):
            name = tag[len(W):]
            if name in _HALF:
                mul(el, "val", minimum=2)
            elif name == "spacing":
                if el.getparent().tag == W + "rPr":
                    mul(el, "val")
                else:
                    mul(el, "before")
                    mul(el, "after")
                    if el.get(W + "lineRule") in ("exact", "atLeast"):
                        mul(el, "line", minimum=20)
            elif name in _TWIPS and name != "spacing":
                if name in ("tblW", "tcW") and el.get(W + "type") not in (None, "dxa"):
                    continue
                for a in _TWIPS[name]:
                    mul(el, a)
            elif name == "tblpPr":
                for a in ("tblpX", "tblpY", "leftFromText", "rightFromText", "topFromText", "bottomFromText"):
                    mul(el, a)
            elif name in ("top", "left", "bottom", "right", "start", "end") and \
                    el.getparent() is not None and el.getparent().tag in (W + "tcMar", W + "tblCellMar"):
                mul(el, "w")
        elif tag.startswith(WP):
            name = tag[len(WP):]
            if name in _EMU:
                for a in _EMU[name]:
                    mul(el, a, ns="")
            elif name == "posOffset" and el.text and el.text.lstrip("-").isdigit():
                rel = el.getparent().get("relativeFrom")
                v = int(el.text)
                if rel == "page" and el.getparent().tag == WP + "positionH":
                    v = int((v / 12700 - old_left) * s * 12700 + LEFT_CM / 2.54 * 72 * 12700)
                else:
                    v = int(v * s)
                el.text = str(v)
        elif tag.startswith(A):
            name = tag[len(A):]
            if name in ("ext", "off", "chExt", "chOff") and el.getparent() is not None \
                    and el.getparent().tag == A + "xfrm":
                for a in ("cx", "cy", "x", "y"):
                    mul(el, a, ns="")


def _empty(el):
    from docx.oxml.ns import qn
    return el.tag == qn("w:p") and not "".join(el.itertext()).strip() \
        and el.find(".//" + qn("w:drawing")) is None and el.find(qn("w:pPr") + "/" + qn("w:sectPr")) is None


def trim_leading(grp, amount):
    """구역 맨 위의 빈 자리(빈 문단 높이·문단 앞 간격)를 amount(pt) 만큼 깎는다."""
    import docx
    from docx.shared import Pt
    from docx.oxml.ns import qn
    for el in list(grp):
        if amount <= 0.5 or el.tag != qn("w:p"):
            break
        pf = docx.text.paragraph.Paragraph(el, None).paragraph_format
        before = pf.space_before.pt if pf.space_before is not None else 0
        cut = min(before, amount)
        pf.space_before = Pt(before - cut)
        amount -= cut
        if not _empty(el):
            break
        line = pf.line_spacing.pt if hasattr(pf.line_spacing, "pt") else 12.0
        after = pf.space_after.pt if pf.space_after is not None else 0
        if line + after <= amount:
            el.getparent().remove(el)
            grp.remove(el)
            amount -= line + after
        else:
            pf.line_spacing = Pt(max(1.0, line - amount))
            amount = 0


def trim_trailing(grp):
    """구역 끝의 빈 문단(pdf2docx 가 쪽 바닥까지 채운 것)을 지우고, 구역 끝 문단은 1pt 로 납작하게.

    ★둘 다 그대로 두면 줄인 쪽 끝에서 다음 쪽으로 넘쳐 **거의 빈 쪽** 이 생겼다(등촌역 4쪽).
      원본 1쪽 = 결과 1쪽이라 다음 구역은 어차피 새 쪽에서 시작하므로 지워도 모양이 같다.
    """
    from docx.oxml.ns import qn
    if not grp:
        return
    last = grp[-1]
    ends = last.tag == qn("w:p") and last.find(qn("w:pPr") + "/" + qn("w:sectPr")) is not None
    for el in reversed(grp[:-1] if ends else list(grp)):
        if _empty(el):
            el.getparent().remove(el)
            grp.remove(el)
        else:
            break
    if ends:
        _flatten_holder(last)


def _flatten_holder(p):
    """구역 끝 문단을 높이 1pt 짜리로(글자 1pt·줄 간격 고정 1pt·앞뒤 0)."""
    from docx.oxml.ns import qn
    from lxml import etree
    for r in list(p.findall(qn("w:r"))):
        if not "".join(r.itertext()).strip():
            p.remove(r)
    ppr = p.find(qn("w:pPr"))
    sp = ppr.find(qn("w:spacing"))
    if sp is None:
        sp = ppr.makeelement(qn("w:spacing"), {})
        sect = ppr.find(qn("w:sectPr"))
        # spacing 은 rPr·sectPr 보다 앞
        anchor = ppr.find(qn("w:jc")) or ppr.find(qn("w:rPr")) or sect
        anchor.addprevious(sp) if anchor is not None else ppr.append(sp)
    for k, v in (("before", "0"), ("after", "0"), ("line", "20"), ("lineRule", "exact")):
        sp.set(qn("w:" + k), v)
    rpr = ppr.find(qn("w:rPr"))
    if rpr is None:
        rpr = ppr.makeelement(qn("w:rPr"), {})
        ppr.find(qn("w:sectPr")).addprevious(rpr)
    for tag in ("w:sz", "w:szCs"):
        e = rpr.find(qn(tag))
        if e is None:
            e = rpr.makeelement(qn(tag), {})
            rpr.append(e)
        e.set(qn("w:val"), "2")


def _insert_picture_page(d, idx, png, box):
    """idx 번째 구역 자리에 원본 쪽 그림 한 장짜리 구역을 넣는다. 쓰인 높이(pt)를 돌려준다."""
    import copy
    from docx.shared import Pt
    from docx.oxml.ns import qn
    s = min(1.0, BOX_W_PT / box.width, BOX_H_PT / box.height)
    w, h = box.width * s, box.height * s
    body = d.element.body
    p = d.add_paragraph()
    p.add_run().add_picture(io.BytesIO(png), width=Pt(w), height=Pt(h))
    holder = d.add_paragraph()
    sp = copy.deepcopy(body.find(qn("w:sectPr")))
    for e in sp.findall(qn("w:headerReference")) + sp.findall(qn("w:footerReference")):
        sp.remove(e)
    holder._p.get_or_add_pPr().append(sp)
    body.remove(p._p)
    body.remove(holder._p)
    groups = [g for g in _section_groups(d) if g]
    if idx < len(groups):
        # idx 번째 구역 앞에 [그림 문단 + 구역 끝] 을 끼운다 → 그림만 든 구역이 하나 생긴다
        groups[idx][0].addprevious(p._p)
        groups[idx][0].addprevious(holder._p)
    else:
        # 맨 끝: 앞 구역을 닫는 '구역 끝' 문단을 먼저 두고, 그림은 마지막 구역(본문 sectPr)에 넣는다
        last = body.find(qn("w:sectPr"))
        holder._p.find(qn("w:pPr")).remove(holder._p.find(qn("w:pPr")).find(qn("w:sectPr")))
        holder._p.get_or_add_pPr().append(copy.deepcopy(last))
        last.addprevious(holder._p)
        last.addprevious(p._p)
    _set_geometry(d.sections[idx])
    return h
