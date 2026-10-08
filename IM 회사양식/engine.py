# -*- coding: utf-8 -*-
"""IM 회사양식 — 만들기 엔진.

원본 IM 워드(.docx)를 바탕으로 두고 회사 양식만 입힌다. 본문 내용·글씨 크기·쪽 방향은 건드리지 않는다.

  1. 글꼴        → Pretendard (보통 글씨 'Pretendard Light', 굵은 글씨 'Pretendard')
                    ※ 'Pretendard Light' 에 굵게를 걸면 워드가 가짜로 두껍게 그려서, 굵은 글씨는 일반 굵기 계열로
  2. 표 머리글   → 네이비(17365D) + 흰 글씨. 나머지 칸은 원본 그대로(붕어빵과 같은 규칙)
  3. 원본 표지   → 뺀다 (첫 '쪽 나눔' 앞까지)
     원본 끝쪽   → 증권사 연락처 쪽이면 뺀다 (마지막 '쪽 나눔' 뒤에 메일·전화번호가 있을 때)
  4. 회사 양식   → 맨 앞에 회사 표지 + 'Ⅰ. 사모사채 개요', 모든 쪽 머리말·꼬리말을 회사 것으로,
                    맨 뒤에 회사 연락처 쪽 + '끝'

양식은 사용자가 만든 `양식\\형식_1.docx` (2026-10-07 확정)에서 가져온다. 웹에는 딜 내용을 지운
`양식\\회사양식_틀.docx` 를 올린다(make_clean_template).
"""
import copy
import io
import os
import re
from datetime import date

from lxml import etree
from PIL import Image

import docx
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml.ns import qn
from docx.shared import Pt

import pages

HERE = os.path.dirname(os.path.abspath(__file__))
# 웹·내 컴퓨터 모두 딜 내용을 지운 틀을 쓴다(같은 결과가 나오게). 형식_1 을 고쳤으면 양식\_틀만들기.py 를 돌릴 것.
TEMPLATE_CANDIDATES = [os.path.join(HERE, "양식", "회사양식_틀.docx"),
                       os.path.join(HERE, "양식", "형식_1.docx")]

NAVY = "17365D"           # 표 머리글 (사용자 양식 연락처 표와 같은 색)
LABEL = "F2F2F2"          # 사모사채 개요 구분 칸
FONT_LIGHT = "Pretendard Light"
FONT_BOLD = "Pretendard"
SYMBOL_FONTS = {"wingdings", "wingdings 2", "wingdings 3", "symbol", "webdings", "marlett",
                "mt extra", "segoe ui symbol", "segoe ui emoji"}

BOND_ROWS = ["사모사채명", "사채유형", "발행인", "기초자산", "발행금액",
             "발행일", "만기일", "금융조건", "이자지급주기"]
COVER_PHOTO_RATIO = 15.92 / 7.85   # 양식 표지 사진 칸 가로:세로

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
NS = {"w": W,
      "wp": "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing",
      "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
      "mc": "http://schemas.openxmlformats.org/markup-compatibility/2006",
      "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
      "v": "urn:schemas-microsoft-com:vml"}
R_EMBED = "{%s}embed" % NS["r"]
R_ID = "{%s}id" % NS["r"]

EMAIL_RE = re.compile(r"[\w.\-]+@[\w\-]+\.[\w.]+")
PHONE_RE = re.compile(r"0\d{1,2}[-)\s.]\d{3,4}[-\s.]\d{4}")


def template_path():
    for p in TEMPLATE_CANDIDATES:
        if os.path.exists(p):
            return p
    return None


# ★워드 서식 항목은 **순서가 정해져 있다.** 순서를 어기면 워드가 '파일이 손상되었습니다' 로 안 연다
#   (실제로 겪음 — rFonts 를 rStyle 앞에, shd 를 vAlign 뒤에 넣었을 때).
_ORDER = {
    "rPr": ["rStyle", "rFonts", "b", "bCs", "i", "iCs", "caps", "smallCaps", "strike", "dstrike",
            "outline", "shadow", "emboss", "imprint", "noProof", "snapToGrid", "vanish", "webHidden",
            "color", "spacing", "w", "kern", "position", "sz", "szCs", "highlight", "u", "effect",
            "bdr", "shd", "fitText", "vertAlign", "rtl", "cs", "em", "lang", "eastAsianLayout",
            "specVanish", "oMath", "rPrChange"],
    "tcPr": ["cnfStyle", "tcW", "gridSpan", "hMerge", "vMerge", "tcBorders", "shd", "noWrap", "tcMar",
             "textDirection", "tcFitText", "vAlign", "hideMark", "headers", "cellIns", "cellDel",
             "cellMerge", "tcPrChange"],
    "tblPr": ["tblStyle", "tblpPr", "tblOverlap", "bidiVisual", "tblStyleRowBandSize",
              "tblStyleColBandSize", "tblW", "jc", "tblCellSpacing", "tblInd", "tblBorders", "shd",
              "tblLayout", "tblCellMar", "tblLook", "tblCaption", "tblDescription", "tblPrChange"],
    "pPr": ["pStyle", "keepNext", "keepLines", "pageBreakBefore", "framePr", "widowControl", "numPr",
            "suppressLineNumbers", "pBdr", "shd", "tabs", "suppressAutoHyphens", "kinsoku", "wordWrap",
            "overflowPunct", "topLinePunct", "autoSpaceDE", "autoSpaceDN", "bidi", "adjustRightInd",
            "snapToGrid", "spacing", "ind", "contextualSpacing", "mirrorIndents", "suppressOverlap", "jc",
            "textDirection", "textAlignment", "textboxTightWrap", "outlineLvl", "divId", "cnfStyle",
            "rPr", "sectPr", "pPrChange"],
}


def _put(parent, child):
    """child 를 parent 안 '제자리' 에 넣는다(이미 같은 것이 있으면 바꿔 끼운다)."""
    order = _ORDER.get(etree.QName(parent).localname)
    name = etree.QName(child).localname
    old = parent.find(child.tag)
    if old is not None:
        parent.replace(old, child)
        return child
    if order and name in order:
        later = set(order[order.index(name) + 1:])
        for i, sib in enumerate(parent):
            if etree.QName(sib).localname in later:
                parent.insert(i, child)
                return child
    parent.append(child)
    return child


def _get(parent, name):
    """parent 의 w:name 자식(없으면 제자리에 새로 만든다)."""
    el = parent.find(qn("w:" + name))
    if el is None:
        el = _put(parent, parent.makeelement(qn("w:" + name), {}))
    return el


# ──────────────────────────────────────────────
# 1. 글꼴
# ──────────────────────────────────────────────
def _is_symbol(rfonts):
    """기호 글꼴인가 — pdf2docx 는 'Wingdings-Regular' 처럼 이름 뒤를 붙여 와서 이름 일부로 본다."""
    for k in ("ascii", "hAnsi", "eastAsia", "cs"):
        v = (rfonts.get(qn("w:" + k)) or "").strip().lower()
        if v and (v in SYMBOL_FONTS or re.search(r"wingding|webding|symbol|marlett", v)):
            return True
    return False


def _set_rfonts(rpr, family):
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = _get(rpr, "rFonts")
    elif _is_symbol(rf):
        return
    for k in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
        rf.attrib.pop(qn("w:" + k), None)
    for k in ("ascii", "hAnsi", "eastAsia", "cs"):
        rf.set(qn("w:" + k), family)


def _on(el, tag):
    """<w:b/> 같은 켜고 끄는 서식이 켜져 있나."""
    if el is None:
        return None
    b = el.find(qn(tag))
    if b is None:
        return None
    return b.get(qn("w:val"), "true").lower() not in ("0", "false", "off")


def _style_bold(styles_el):
    """스타일 id → 굵게 여부(바탕 스타일까지 따라감)."""
    raw = {}
    for s in styles_el.findall(qn("w:style")):
        sid = s.get(qn("w:styleId"))
        based = s.find(qn("w:basedOn"))
        raw[sid] = (_on(s.find(qn("w:rPr")), "w:b"), based.get(qn("w:val")) if based is not None else None)
    out = {}
    for sid in raw:
        seen, cur, val = set(), sid, None
        while cur in raw and cur not in seen and val is None:
            seen.add(cur)
            val, cur = raw[cur][0], raw[cur][1]
        out[sid] = bool(val)
    return out


def apply_fonts(doc):
    styles_el = doc.styles.element
    bold_of = _style_bold(styles_el)
    # 기본값·스타일
    for rpr in styles_el.iter(qn("w:rPr")):
        st = rpr.getparent()
        while st is not None and st.tag != qn("w:style"):
            st = st.getparent()
        bold = bold_of.get(st.get(qn("w:styleId")), False) if st is not None else False
        _set_rfonts(rpr, FONT_BOLD if bold else FONT_LIGHT)
    dd = styles_el.find(qn("w:docDefaults"))
    if dd is not None:
        rprd = dd.find(qn("w:rPrDefault"))
        if rprd is None:
            rprd = dd.makeelement(qn("w:rPrDefault"), {})
            dd.insert(0, rprd)                     # rPrDefault 는 pPrDefault 앞
        rpr = rprd.find(qn("w:rPr"))
        if rpr is None:
            rpr = etree.SubElement(rprd, qn("w:rPr"))
        _set_rfonts(rpr, FONT_LIGHT)
    # 테마 글꼴(스타일이 테마를 가리키는 문서가 있다)
    for rel in doc.part.rels.values():
        if rel.reltype == RT.THEME:
            part = rel.target_part
            blob = part.blob.decode("utf-8")
            blob = re.sub(r'(<a:(?:latin|ea) typeface=")[^"]*(")', r"\g<1>%s\2" % FONT_LIGHT, blob)
            blob = re.sub(r'(<a:font script="Hang" typeface=")[^"]*(")', r"\g<1>%s\2" % FONT_LIGHT, blob)
            part._blob = blob.encode("utf-8")
    # 본문의 글자 하나하나(글상자 안까지)
    body = doc.element.body
    for r in body.iter(qn("w:r")):
        rpr = r.find(qn("w:rPr"))
        if rpr is None:
            continue
        p = r.getparent()
        while p is not None and p.tag != qn("w:p"):
            p = p.getparent()
        ps = None
        if p is not None:
            ppr = p.find(qn("w:pPr"))
            if ppr is not None and ppr.find(qn("w:pStyle")) is not None:
                ps = ppr.find(qn("w:pStyle")).get(qn("w:val"))
        rs = rpr.find(qn("w:rStyle"))
        direct = _on(rpr, "w:b")
        bold = direct if direct is not None else (
            bold_of.get(rs.get(qn("w:val")), False) if rs is not None else bold_of.get(ps, False))
        _set_rfonts(rpr, FONT_BOLD if bold else FONT_LIGHT)
    # 번호·글머리표(기호 글꼴은 그대로)
    try:
        num = doc.part.numbering_part.element
        for rpr in num.iter(qn("w:rPr")):
            _set_rfonts(rpr, FONT_LIGHT)
    except Exception:
        pass


# ──────────────────────────────────────────────
# 2. 표 머리글 → 네이비
# ──────────────────────────────────────────────
def _fill(tc):
    tcpr = tc.find(qn("w:tcPr"))
    shd = tcpr.find(qn("w:shd")) if tcpr is not None else None
    if shd is None:
        return None
    f = (shd.get(qn("w:fill")) or "").upper()
    return None if f in ("", "AUTO", "FFFFFF") else f


def _lum(hexs):
    try:
        r, g, b = (int(hexs[i:i + 2], 16) / 255 for i in (0, 2, 4))
    except Exception:
        return 1.0
    return 0.299 * r + 0.587 * g + 0.114 * b


def _first_child(parent, name):
    """parent 맨 앞에 w:name 자식(pPr·rPr·tcPr 처럼 늘 맨 앞인 것)."""
    el = parent.find(qn("w:" + name))
    if el is None:
        el = parent.makeelement(qn("w:" + name), {})
        parent.insert(0, el)
    return el


def _paint_navy(tc):
    tcpr = _first_child(tc, "tcPr")
    shd = tcpr.makeelement(qn("w:shd"), {qn("w:val"): "clear", qn("w:color"): "auto", qn("w:fill"): NAVY})
    _put(tcpr, shd)
    # 글씨는 흰색 — 빈 칸에도 박아 둔다(나중에 타이핑하면 흰 글씨로 나오게)
    for p in tc.findall(qn("w:p")):
        ppr = _first_child(p, "pPr")
        targets = [_get(ppr, "rPr")] + [_first_child(r, "rPr") for r in p.findall(qn("w:r"))]
        for rpr in targets:
            _put(rpr, rpr.makeelement(qn("w:color"), {qn("w:val"): "FFFFFF"}))


def _is_frame(tbl):
    """'큰 틀' 표인가 — 맨 윗줄이 표 폭 전체를 덮는 진한 제목 띠 하나.

    예) 안산 '정보통신업 사업자 관련 설명' : 남색 제목 띠 아래에 작은 표 두 개(회색 머리글)가 들어 있다.
        pdf2docx 는 이걸 표 하나로 합쳐 만들어서, 0번 줄 = 큰 틀 제목, 1번 줄 = 안쪽 표 머리글이 된다.
    """
    rows = tbl.findall(qn("w:tr"))
    if len(rows) < 2:
        return False
    first = rows[0].findall(qn("w:tc"))
    ncols = len(tbl.findall(qn("w:tblGrid") + "/" + qn("w:gridCol"))) or max(len(r.findall(qn("w:tc"))) for r in rows)
    if len(first) == 1:
        span = first[0].find(qn("w:tcPr") + "/" + qn("w:gridSpan"))
        wide = ncols <= 1 or (span is not None and int(span.get(qn("w:val"), "1")) >= ncols - 1)
    else:
        wide = False
    f = _fill(first[0]) if first else None
    return wide and bool(f) and _lum(f) < 0.5


def paint_headers(doc):
    """머리글 칸을 회사 네이비로. 고친 칸 수를 돌려준다.

    규칙(사용자 확정)
      · 보통 표 : 맨 위 1~2줄 머리글은 색이 연해도 네이비 + 흰 글씨(붕어빵과 같다), 진한 칸은 어디 있든 네이비.
      · 표 안의 표('큰 틀') : **큰 틀 제목 띠만** 회사 네이비. 그 안 작은 표들의 색은 **원본 그대로**(2026-10-08).
        원본이 이미 네이비여도 회사 네이비(17365D)로 맞춘다.
    """
    n = 0
    body = doc.element.body
    frames = [t for t in body.iter(qn("w:tbl")) if _is_frame(t)]
    frame_ids = set(map(id, frames))
    for tbl in body.iter(qn("w:tbl")):
        if any(a in frames for a in tbl.iterancestors(qn("w:tbl"))):
            continue                                          # 큰 틀 안의 표 — 원본 색 그대로
        rows = tbl.findall(qn("w:tr"))
        if not rows:
            continue
        cells = [r.findall(qn("w:tc")) for r in rows]
        if tbl in frames:
            _paint_navy(cells[0][0])
            n += 1
            continue
        if len(rows) < 2 or max(len(c) for c in cells) < 2:
            for row in cells:                                 # 한 줄짜리 제목 띠 등 — 진한 칸만
                for tc in row:
                    f = _fill(tc)
                    if f and _lum(f) < 0.5 and f != NAVY:
                        _paint_navy(tc)
                        n += 1
            continue
        head = set()
        for i, r in enumerate(rows):                      # '머리글 줄 반복' 표시가 된 줄
            trpr = r.find(qn("w:trPr"))
            if trpr is not None and _on(trpr, "w:tblHeader"):
                head.add(i)
        body_fills = [_fill(tc) for row in cells[1:] for tc in row]
        common = max(set(body_fills), key=body_fills.count) if body_fills else None
        for i in (0, 1):                                  # 맨 위 1~2줄이 칠해져 있으면 머리글
            if i >= len(rows) - 1:
                break
            fills = [_fill(tc) for tc in cells[i]]
            colored = [f for f in fills if f]
            if len(colored) >= max(1, 0.6 * len(fills)) and (i == 0 or i - 1 in head) \
                    and max(set(colored), key=colored.count) != common:
                head.add(i)
            else:
                break
        for i, row in enumerate(cells):
            for tc in row:
                f = _fill(tc)
                if i in head or (f and _lum(f) < 0.5):     # 진한 칸은 어디 있든 머리글로 본다
                    if f != NAVY or i in head:
                        _paint_navy(tc)
                        n += 1
    return n


# ──────────────────────────────────────────────
# 3. 원본 표지·끝쪽
# ──────────────────────────────────────────────
def _hard_break(el):
    """이 요소가 '여기서 쪽이 끝난다' 를 품고 있나 → 'run'(쪽나눔) / 'sect'(구역나눔) / None"""
    if el.tag != qn("w:p"):
        return None
    ppr = el.find(qn("w:pPr"))
    if ppr is not None:
        sp = ppr.find(qn("w:sectPr"))
        if sp is not None:
            t = sp.find(qn("w:type"))
            if t is None or t.get(qn("w:val")) != "continuous":
                return "sect"
    for br in el.iter(qn("w:br")):
        if br.get(qn("w:type")) == "page":
            return "run"
    return None


def _starts_page(el):
    ppr = el.find(qn("w:pPr")) if el.tag == qn("w:p") else None
    return ppr is not None and _on(ppr, "w:pageBreakBefore")


def _text(els):
    return "".join("".join(e.itertext()) for e in els)


def _blocks(body):
    return [e for e in body if e.tag in (qn("w:p"), qn("w:tbl"), qn("w:sdt"))]


def find_cover(doc):
    """원본 표지 = (끝 블록 번호, 근거, 표지 글). 못 찾으면 (None, 이유, '')."""
    end, why = pages.cover_end(doc)
    if end is None:
        return None, why, ""
    bl = _blocks(doc.element.body)
    return end, why, " ".join(pages.text_of(e).strip() for e in bl[:end + 1] if pages.text_of(e).strip())


def find_contact_tail(doc):
    """맨 끝 쪽이 증권사 연락처 쪽이면 (시작 블록 번호, 글, 연락처 수). 아니면 (None, 글, 수)."""
    start = pages.last_page_start(doc)
    if start is None:
        return None, "", 0
    bl = _blocks(doc.element.body)
    txt = " ".join(pages.text_of(e) for e in bl[start:])
    hits = len(set(EMAIL_RE.findall(txt))) + len(set(PHONE_RE.findall(txt)))
    if hits >= 2 and len(txt) < 3000:
        return start, txt, hits
    return None, txt, hits


def remove_cover(doc, end):
    """블록 0 ~ end 를 뺀다. 표지 끝이 구역 나눔이면 그 구역 설정도 함께 사라진다(표지 전용 설정)."""
    if end is None:
        return False
    bl = _blocks(doc.element.body)
    for el in bl[:end + 1]:
        el.getparent().remove(el)
    return True


def remove_tail(doc, start):
    body = doc.element.body
    bl = _blocks(body)
    prev = bl[start - 1] if start > 0 else None
    for el in bl[start:]:
        body.remove(el)
    if prev is not None and _hard_break(prev) == "sect":
        # 마지막 구역의 쪽 설정이 사라졌다 → 앞 구역 설정을 문서 끝 설정으로 올린다
        sp = prev.find(qn("w:pPr")).find(qn("w:sectPr"))
        prev.find(qn("w:pPr")).remove(sp)
        old = body.find(qn("w:sectPr"))
        body.replace(old, sp)


def strip_trailing_break(doc):
    """본문 맨 끝의 쪽나눔(뒤에 우리 연락처 쪽이 붙으면 빈 쪽이 생긴다)."""
    bl = _blocks(doc.element.body)
    for el in reversed(bl):
        if el.tag != qn("w:p"):
            return
        brs = [b for b in el.iter(qn("w:br")) if b.get(qn("w:type")) == "page"]
        if brs:
            for b in brs:
                r = b.getparent()
                r.remove(b)
            return
        if "".join(el.itertext()).strip() or el.find(".//" + qn("w:drawing")) is not None:
            return


# ──────────────────────────────────────────────
# 원본에서 표지 제목·사진 후보 뽑기
# ──────────────────────────────────────────────
def suggest_title(doc, end=None):
    """원본 표지에서 가장 큰 글씨 줄들 → 표지 제목 후보(최대 2줄)."""
    if end is None:
        end, _, _ = find_cover(doc)
    bl = _blocks(doc.element.body)[: (end + 1) if end is not None else 10]
    lines = []
    for el in bl:
        for p in el.iter(qn("w:p")):
            if any(a.tag == pages.FB for a in p.iterancestors()):
                continue
            if p.find(".//" + qn("w:txbxContent")) is not None:
                continue                                   # 글상자를 품은 문단은 안쪽 문단에서 센다
            t = re.sub(r"\s+", " ", pages.text_of(p)).strip()
            if not t or len(t) > 60:
                continue
            sizes = [int(s.get(qn("w:val"))) for s in p.iter(qn("w:sz")) if (s.get(qn("w:val")) or "").isdigit()]
            lines.append((max(sizes) if sizes else 20, t))
    lines = [(s, t) for s, t in lines
             if not re.search(r"memorandum|confidential|private|^\d{4}[.\-/]?\d{0,2}\.?$", t, re.I)]
    if not lines:
        return []
    top = max(s for s, _ in lines)
    out = []
    for s, t in lines:
        if s >= top - 4 and t not in out:
            out.append(t)
    return out[:2]


def photo_candidates(doc, limit=6):
    """본문 그림 중 표지 사진으로 쓸 만한 것(큰 실사 사진 먼저)."""
    seen, out = set(), []
    for rel in doc.part.rels.values():
        if rel.reltype != RT.IMAGE or rel.is_external:
            continue
        part = rel.target_part
        if part.partname in seen:
            continue
        seen.add(part.partname)
        try:
            im = Image.open(io.BytesIO(part.blob))
            w, h = im.size
            if w < 300 or h < 150:
                continue
            small = im.convert("L").resize((60, 60))
            bright = sum(small.getdata()) / 3600
            # 실사 사진은 색이 수천 가지, 구조도·표 그림은 몇십 가지 — 어두운 구조도를 사진으로
            # 고르던 것을 막는다(실제: 엘피스 표지에 금융구조도가 골라졌다)
            colors = len(set(im.convert("RGB").resize((80, 80)).getdata()))
        except Exception:
            continue
        ratio = w / h
        score = (colors > 1500, bright < 210, 1.2 <= ratio <= 2.6, w * h)
        out.append((score, part.blob))
    out.sort(key=lambda t: t[0], reverse=True)
    return [b for _, b in out[:limit]]


# ──────────────────────────────────────────────
# 4. 회사 양식 붙이기
# ──────────────────────────────────────────────
def _prepare_photo(blob):
    """표지 사진: **자르지 않는다.** 원본 그대로 두고 화질만 또렷하게(2026-10-08 사용자 확정).

    · 가로 1,800px 보다 작으면 같은 비율로 키운 뒤(인쇄 때 뭉개지지 않게) 윤곽만 살짝 또렷하게 한다.
    · 다시 저장할 때 화질이 떨어지지 않게 PNG(투명 있으면) 또는 JPEG 95 로 저장.
    반환: (bytes, (가로px, 세로px))
    """
    from PIL import ImageFilter
    im = Image.open(io.BytesIO(blob))
    has_alpha = im.mode in ("RGBA", "LA") or (im.mode == "P" and "transparency" in im.info)
    im = im.convert("RGBA" if has_alpha else "RGB")
    if im.width < 1800:
        k = 1800 / im.width
        im = im.resize((1800, max(1, int(im.height * k))), Image.LANCZOS)
    im = im.filter(ImageFilter.UnsharpMask(radius=1.2, percent=70, threshold=2))
    out = io.BytesIO()
    if has_alpha:
        im.save(out, "PNG")
    else:
        im.save(out, "JPEG", quality=95, subsampling=0)
    return out.getvalue(), im.size


def _anchor_of(node):
    while node is not None and etree.QName(node).localname not in ("anchor", "inline"):
        node = node.getparent()
    return node


def _pos_v(anchor):
    po = anchor.find("wp:positionV/wp:posOffset", NS)
    return po


def _fit_cover_photo(blip_node, pw, ph):
    """사진을 원본 비율 그대로 '사진 칸' 안에 들어가게 줄이고 가로 가운데에 둔다.

    사진 칸 = 양식 사진 가로 폭 × (사진 윗자리 ~ Disclaimer 상자 위 0.4cm).
    """
    anc = _anchor_of(blip_node)
    if anc is None:
        return
    ext = anc.find("wp:extent", NS)
    max_w = int(ext.get("cx"))
    top = _pos_v(anc)
    room = int(anc.get("_room_cy") or ext.get("cy"))
    w = max_w
    h = int(w * ph / pw)
    if h > room:
        h = room
        w = int(h * pw / ph)
    ext.set("cx", str(w))
    ext.set("cy", str(h))
    for x in anc.iter("{%s}ext" % NS["a"]):
        if x.getparent() is not None and etree.QName(x.getparent()).localname == "xfrm":
            x.set("cx", str(w))
            x.set("cy", str(h))
    ph_el = anc.find("wp:positionH", NS)
    off = ph_el.find("wp:posOffset", NS) if ph_el is not None else None
    if off is not None:
        off.text = str(int(off.text) + (max_w - w) // 2)
    anc.attrib.pop("_room_cy", None)


def _text_width_pt(text, size_pt, bold=True):
    """Pretendard 로 찍었을 때 글 폭(pt) 어림 — 웹에는 글꼴 파일이 없어 글자 종류별 평균 폭으로 잰다."""
    w = 0.0
    for ch in text:
        if ord(ch) >= 0x1100:                 # 한글·한자
            w += 0.92
        elif ch == " ":
            w += 0.28
        elif ch.isupper() or ch.isdigit():
            w += 0.64
        else:
            w += 0.55
    return w * size_pt * (1.04 if bold else 1.0)


def _layout_cover(els, lines):
    """제목 상자를 글 길이에 맞춰 늘리고, 그만큼 날짜·사진을 아래로 내린다(2026-10-08 사용자 지적).

    양식 제목 상자는 두 줄 높이다. 둘째 줄이 길어 상자 폭을 넘으면 세 줄 이상이 되는데, 상자가 그대로면
    아래 날짜와 겹친다(실제: '안산 OUR IDC PF 대출 / 부산_남포동_생활형숙박시설 신축사업 PF대출 후순위').
    """
    anchors = [a for el in els for a in el.iter("{%s}anchor" % NS["wp"])]
    title = date_a = pic = disc = None
    for a in anchors:
        t = "".join(x.text or "" for x in a.iter(qn("w:t")))
        if a.find(".//a:blip", NS) is not None:
            pic = a
        elif "Disclaimer" in t:
            disc = a
        elif re.search(r"(January|February|March|April|May|June|July|August|September|October|November|December)\s+\d{4}", t):
            date_a = a
        elif a.find(".//" + qn("w:txbxContent")) is not None and title is None and \
                any(l and l in t for l in lines if l):
            title = a
    if title is None:
        return
    ext = title.find("wp:extent", NS)
    body_pr = title.find(".//{http://schemas.microsoft.com/office/word/2010/wordprocessingShape}bodyPr")
    l_ins = int(body_pr.get("lIns", 91440)) if body_pr is not None else 91440
    r_ins = int(body_pr.get("rIns", 91440)) if body_pr is not None else 91440
    width_pt = (int(ext.get("cx")) - l_ins - r_ins) / 12700
    paras = title.find(".//" + qn("w:txbxContent")).findall(qn("w:p"))
    need = 0
    for p in paras:
        txt = "".join(x.text or "" for x in p.iter(qn("w:t")))
        sizes = [int(s.get(qn("w:val"))) / 2 for s in p.iter(qn("w:sz")) if (s.get(qn("w:val")) or "").isdigit()]
        size = max(sizes) if sizes else 20
        if txt.strip():
            need += max(1, -(-int(_text_width_pt(txt, size)) // max(int(width_pt), 1)))
    have = sum(1 for p in paras if "".join(x.text or "" for x in p.iter(qn("w:t"))).strip()) or 1
    per_line = int(ext.get("cy")) // max(len(paras), 1)
    extra = max(0, need - max(have, len(paras))) * per_line
    if extra:
        ext.set("cy", str(int(ext.get("cy")) + extra))
        for x in title.iter("{%s}ext" % NS["a"]):
            if x.getparent() is not None and etree.QName(x.getparent()).localname == "xfrm":
                x.set("cy", str(int(x.get("cy")) + extra))
        for a in (date_a, pic):
            if a is not None and _pos_v(a) is not None:
                _pos_v(a).text = str(int(_pos_v(a).text) + extra)
    if pic is not None and disc is not None and _pos_v(pic) is not None and _pos_v(disc) is not None:
        room = int(_pos_v(disc).text) - int(_pos_v(pic).text) - 144000      # Disclaimer 위 0.4cm 띄움
        pic.set("_room_cy", str(max(room, 360000)))


def _set_par_text(p, text):
    """문단 서식은 두고 글만 바꾼다(첫 글자 조각에 몰아 넣음)."""
    ts = list(p.iter(qn("w:t")))
    if not ts:
        return
    ts[0].text = text
    ts[0].set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    for t in ts[1:]:
        t.text = ""


def _cover_elements(tpl, doc, title_lines, date_text, photo):
    """양식 표지(문단 0·1)를 복사해 글·사진을 바꿔 돌려준다."""
    tb = _blocks(tpl.element.body)
    els = [copy.deepcopy(tb[0]), copy.deepcopy(tb[1])]
    lines = (list(title_lines) + ["", ""])[:2]
    for el in els:
        for tx in el.iter(qn("w:txbxContent")):
            t = "".join(tx.itertext())
            ps = tx.findall(qn("w:p"))
            if "제안서" in t and len(ps) >= 2:            # 제목 상자(두 줄)
                _set_par_text(ps[0], lines[0])
                _set_par_text(ps[1], lines[1])
            elif re.search(r"(January|February|March|April|May|June|July|August|September|"
                           r"October|November|December)\s+\d{4}", t):
                _set_par_text(ps[0], date_text)
    # 글상자의 '대체 내용'(옛 워드용 VML)은 뺀다 — 상자 크기·자리를 바꿀 때 두 벌을 맞출 필요가 없게.
    # 지금 워드는 mc:Choice 쪽만 그린다.
    for el in els:
        for fb in list(el.iter("{%s}Fallback" % NS["mc"])):
            fb.getparent().remove(fb)
    _layout_cover(els, lines)
    # 사진: 원본 rId 를 새 문서의 그림으로 바꾸거나, 사진 없음이면 그림을 뺀다
    tpl_rels = tpl.part.rels
    for el in els:
        for node in list(el.iter()):
            for attr in (R_EMBED, R_ID):
                rid = node.get(attr)
                if not rid or rid not in tpl_rels or tpl_rels[rid].reltype != RT.IMAGE:
                    continue
                if photo is None:
                    box = node
                    while box is not None and etree.QName(box).localname not in ("drawing", "pict"):
                        box = box.getparent()
                    if box is not None and box.getparent() is not None:
                        box.getparent().remove(box)
                    break
                blob, (pw, ph) = _prepare_photo(photo)
                new_rid, _ = doc.part.get_or_add_image(io.BytesIO(blob))
                node.set(attr, new_rid)
                _fit_cover_photo(node, pw, ph)
    for el in els:
        _import_styles(tpl, doc, el)
        _pin_template_defaults(doc, el)
    return els


def _pin_template_defaults(doc, root):
    """양식에서 가져온 부분에 양식의 '기본값' 을 박는다(글자 10pt, 문단 앞뒤 0, 줄간격 1).

    ★양식은 기본값(글자 크기·문단 간격)을 따로 안 적고 문서 기본값에 기대는데, 원본 문서의 기본값은
      다르다. pdf2docx 문서는 기본 11pt·문단 뒤 10pt 라 표지 'Strictly Confidential' 상자가 넘쳐
      'Strictly' 만 보였다(2026-10-07). 직접 적었거나 회사양식 서식(cf_)이 정한 값은 그대로 둔다.
    """
    st = {s.get(qn("w:styleId")): s for s in doc.styles.element.findall(qn("w:style"))}

    def style_has(sid, path):
        seen = set()
        while sid in st and sid not in seen:
            seen.add(sid)
            if st[sid].find(path) is not None:
                return True
            b = st[sid].find(qn("w:basedOn"))
            sid = b.get(qn("w:val")) if b is not None else None
        return False

    for p in root.iter(qn("w:p")):
        ppr = _first_child(p, "pPr")
        ps = ppr.find(qn("w:pStyle"))
        sid = ps.get(qn("w:val")) if ps is not None else None
        if ppr.find(qn("w:spacing")) is None and not style_has(sid, qn("w:pPr") + "/" + qn("w:spacing")):
            _put(ppr, ppr.makeelement(qn("w:spacing"), {qn("w:before"): "0", qn("w:after"): "0",
                                                         qn("w:line"): "240", qn("w:lineRule"): "auto"}))
        for r in p.findall(qn("w:r")) + [ppr]:
            rpr = _get(ppr, "rPr") if r is ppr else _first_child(r, "rPr")
            rs = rpr.find(qn("w:rStyle"))
            if rpr.find(qn("w:sz")) is None and not style_has(sid, qn("w:rPr") + "/" + qn("w:sz")) \
                    and not (rs is not None and style_has(rs.get(qn("w:val")), qn("w:rPr") + "/" + qn("w:sz"))):
                _put(rpr, rpr.makeelement(qn("w:sz"), {qn("w:val"): "20"}))


def _import_styles(tpl, doc, root):
    """root 안에서 쓰는 양식 서식을 원본에 'cf_<이름>' 으로 복사해 붙이고, 참조를 그쪽으로 바꾼다.

    ★서식 이름(styleId)이 같아도 문서마다 뜻이 다르다. 안산 원본에서는 양식의 표지 서식 이름이
      '제N장' 번호 서식이라 표지에 '제1장…제6장' 이, 머리말 앞에 '1)' 이 붙었다(2026-10-07).
    """
    tst = tpl.styles.element
    dst = doc.styles.element
    by_id = {s.get(qn("w:styleId")): s for s in tst.findall(qn("w:style"))}
    have = {s.get(qn("w:styleId")) for s in dst.findall(qn("w:style"))}

    def bring(sid):
        if sid not in by_id:
            return None
        nid = "cf_" + sid
        if nid in have:
            return nid
        have.add(nid)
        st = copy.deepcopy(by_id[sid])
        st.set(qn("w:styleId"), nid)
        st.attrib.pop(qn("w:default"), None)
        nm = st.find(qn("w:name"))
        if nm is not None:
            nm.set(qn("w:val"), "회사양식 " + nm.get(qn("w:val")))
        for numpr in list(st.iter(qn("w:numPr"))):         # 번호 매기기는 가져오지 않는다
            numpr.getparent().remove(numpr)
        for tag in ("w:basedOn", "w:next", "w:link"):
            ref = st.find(qn(tag))
            if ref is not None:
                new = bring(ref.get(qn("w:val")))
                if new:
                    ref.set(qn("w:val"), new)
                else:
                    st.remove(ref)
        dst.append(st)
        return nid

    for tag in ("w:pStyle", "w:rStyle", "w:tblStyle"):
        for s in list(root.iter(qn(tag))):
            new = bring(s.get(qn("w:val")))
            if new:
                s.set(qn("w:val"), new)
            else:
                s.getparent().remove(s)


def _run(p, text, size, bold=False, color=None, font=None):
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.bold = bold
    rpr = r._r.get_or_add_rPr()
    _set_rfonts(rpr, font or (FONT_BOLD if bold else FONT_LIGHT))
    if color:
        from docx.shared import RGBColor
        r.font.color.rgb = RGBColor.from_string(color)
    return r


def _para_fmt(p, align="left", before=0, after=0):
    from docx.enum.text import WD_ALIGN_PARAGRAPH as A
    p.alignment = {"left": A.LEFT, "center": A.CENTER}[align]
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)


def _cell_shade(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    _put(tcpr, tcpr.makeelement(qn("w:shd"), {qn("w:val"): "clear", qn("w:color"): "auto", qn("w:fill"): fill}))


def _table_borders(tbl, color="000000", sz=4):
    tblpr = tbl._tbl.tblPr
    b = _put(tblpr, tblpr.makeelement(qn("w:tblBorders"), {}))
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = etree.SubElement(b, qn("w:" + side))
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), str(sz))
        e.set(qn("w:space"), "0")
        e.set(qn("w:color"), color)


def _fixed_widths(tbl, widths_twips):
    from docx.shared import Twips
    tbl.autofit = False
    for row in tbl.rows:
        for c, w in zip(row.cells, widths_twips):
            c.width = Twips(w)
    grid = tbl._tbl.tblGrid
    for gc, w in zip(grid.findall(qn("w:gridCol")), widths_twips):
        gc.set(qn("w:w"), str(w))


def _row_height(row, twips):
    from docx.shared import Twips
    from docx.enum.table import WD_ROW_HEIGHT_RULE
    row.height = Twips(twips)
    row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST


def _vcenter(cell):
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT as V
    cell.vertical_alignment = V.CENTER


def _build_bond_section(doc, bond):
    """'Ⅰ. 사모사채 개요' 제목 + 9줄 표. 새로 만든 요소들을 돌려준다(본문 끝에 붙었다가 떼어 냄)."""
    body = doc.element.body
    # ★id() 로 가르면 안 된다 — lxml 은 파이썬 쪽 객체를 버렸다 다시 만들어 번호가 재사용된다
    #   (그래서 연락처 표가 2쪽에, 개요 제목이 맨 끝에 가 있었다). 목록을 붙들고 'is' 로 가른다.
    before = list(body)
    h = doc.add_paragraph()
    _para_fmt(h, "left", 0, 10)
    _run(h, "Ⅰ. 사모사채 개요", 16, bold=True)
    t = doc.add_table(rows=len(BOND_ROWS), cols=2)
    _table_borders(t)
    _fixed_widths(t, [1448, 7619])
    for i, key in enumerate(BOND_ROWS):
        a, b = t.rows[i].cells
        _row_height(t.rows[i], 454)
        _cell_shade(a, LABEL)
        for c in (a, b):
            _vcenter(c)
        pa, pb = a.paragraphs[0], b.paragraphs[0]
        _para_fmt(pa, "center")
        _para_fmt(pb, "left")
        _run(pa, "사채 유형" if key == "사채유형" else key, 10)
        _run(pb, (bond.get(key) or "").strip(), 10)
    return [e for e in body if not any(e is b for b in before) and e.tag != qn("w:sectPr")]


def _build_contact_section(doc, contacts):
    body = doc.element.body
    # ★id() 로 가르면 안 된다 — lxml 은 파이썬 쪽 객체를 버렸다 다시 만들어 번호가 재사용된다
    #   (그래서 연락처 표가 2쪽에, 개요 제목이 맨 끝에 가 있었다). 목록을 붙들고 'is' 로 가른다.
    before = list(body)
    cols = ["본부", "직급", "이름", "E-mail", "연락처"]
    t = doc.add_table(rows=1 + len(contacts), cols=len(cols))
    _table_borders(t)
    _fixed_widths(t, [2302, 1170, 1459, 2182, 1903])
    for j, name in enumerate(cols):
        c = t.rows[0].cells[j]
        _cell_shade(c, NAVY)
        _vcenter(c)
        p = c.paragraphs[0]
        _para_fmt(p, "center")
        _run(p, name, 9, bold=True, color="FFFFFF")
    _row_height(t.rows[0], 397)
    for i, row in enumerate(contacts, start=1):
        _row_height(t.rows[i], 397)
        for j, key in enumerate(cols):
            c = t.rows[i].cells[j]
            _vcenter(c)
            p = c.paragraphs[0]
            _para_fmt(p, "center")
            _run(p, row.get(key, ""), 9)
    p = doc.add_paragraph()
    _para_fmt(p, "left", 8, 0)
    _run(p, "본 자료와 관련하여 문의사항이 있으실 경우 아래의 담당자에게 언제든지 연락하여 주시면 "
            "성심껏 답변하여 드리겠습니다.", 9)
    p = doc.add_paragraph()
    _para_fmt(p, "center", 220, 0)
    _run(p, "끝", 11)
    return [e for e in body if not any(e is b for b in before) and e.tag != qn("w:sectPr")]


def read_contacts(tpl):
    """양식 마지막 쪽 연락처 표에서 담당자 목록을 읽는다(빈 칸은 건너뜀)."""
    for t in tpl.tables:
        head = [c.text.strip() for c in t.rows[0].cells]
        if "이름" in head and "연락처" in head:
            out, cols = [], ["본부", "직급", "이름", "E-mail", "연락처"]
            for r in t.rows[1:]:
                vals = []
                seen = set()
                for c in r.cells:
                    if id(c._tc) in seen:
                        continue
                    seen.add(id(c._tc))
                    vals.append(c.text.strip())
                vals = [v for v in vals]
                # 사용자 양식에는 빈 칸이 하나 끼어 있다(6칸) → 머리글 이름으로 맞춘다
                cells = [c.text.strip() for c in r.cells]
                row = {}
                for k in cols:
                    j = head.index(k)
                    row[k] = cells[j]
                if not row["본부"]:
                    row["본부"] = next((v for v in cells[:head.index("직급")] if v), "")
                out.append(row)
            return out
    return []


def _copy_hdrftr(tpl, doc, src_part, dst_hf, header_text=None):
    """양식 머리말/꼬리말 내용을 이 문서의 머리말/꼬리말로 옮긴다(그림 연결까지)."""
    # ★조각만 옮기면 안 된다 — 글상자의 mc:Choice Requires="wps" 가 가리키는 이름표(xmlns:wps)가
    #   맨 위(w:hdr)에 선언돼 있어야 워드가 연다. 그래서 양식 머리말 전체를 선언째 통째로 바꿔 끼운다.
    #   (조각만 옮겼을 때 워드가 '파일이 손상되었습니다' 로 안 열었다 — 2026-10-07)
    part = dst_hf.part
    part._element = copy.deepcopy(src_part.element)
    dst = part._element
    _import_styles(tpl, doc, dst)
    _pin_template_defaults(doc, dst)
    for node in dst.iter():
        for attr in (R_EMBED, R_ID):
            rid = node.get(attr)
            if rid and rid in src_part.rels and src_part.rels[rid].reltype == RT.IMAGE:
                blob = src_part.rels[rid].target_part.blob
                new_rid, _ = dst_hf.part.get_or_add_image(io.BytesIO(blob))
                node.set(attr, new_rid)
    if header_text is not None:
        for p in dst.iter(qn("w:p")):
            if "".join(p.itertext()).strip():
                _set_par_text(p, header_text)
                break
    for rpr in dst.iter(qn("w:rPr")):
        rf = rpr.find(qn("w:rFonts"))
        if rf is not None and not _is_symbol(rf):
            pass


def _renumber_shapes(doc):
    """그림·도형 번호(wp:docPr id)를 문서 전체에서 겹치지 않게 다시 매긴다.

    ★양식에서 옮겨 온 표지·꼬리말 도형 번호가 원본 그림 번호와 겹치면 워드가
      '파일이 손상되었습니다' 로 안 연다(실제로 겪음 — 꼬리말 도형 3·4번이 원본과 겹침).
    """
    n = 0
    for part in doc.part.package.iter_parts():
        el = getattr(part, "_element", None)
        if el is None:
            continue
        for dp in el.iter("{%s}docPr" % NS["wp"]):
            n += 1
            dp.set("id", str(n))


def _merge_nsmap(doc, tpl):
    """양식 표지 글상자가 쓰는 이름표(wps 등) 선언을 원본 문서 맨 위에 합친다.

    mc:Choice Requires="wps" 는 맨 위에 xmlns:wps 가 있어야 워드가 연다. 워드로 만든 문서는 대개
    있지만, 다른 프로그램(pdf2docx 등)이 만든 문서에는 없을 수 있다.
    """
    root = doc.part._element
    extra = {k: v for k, v in tpl.part.element.nsmap.items() if k and k not in root.nsmap}
    if not extra:
        return
    # python-docx 의 파서로 만들어야 '문서' 요소(body 등)로 동작한다 — 그냥 lxml 로 만들면 기능이 없다
    from docx.oxml.parser import oxml_parser
    new = oxml_parser.makeelement(root.tag, nsmap={**root.nsmap, **extra})
    for k, v in root.attrib.items():
        new.set(k, v)
    for ch in list(root):
        new.append(ch)
    doc.part._element = new
    doc._element = new
    doc._Document__body = None          # python-docx 가 붙들고 있던 본문 객체를 새 뿌리로 다시 잡게


def _sect_from_template(tpl, title_page):
    sp = copy.deepcopy(tpl.element.body.find(qn("w:sectPr")))
    for tag in ("w:headerReference", "w:footerReference"):
        for e in sp.findall(qn(tag)):
            sp.remove(e)
    tp = sp.find(qn("w:titlePg"))
    if not title_page and tp is not None:
        sp.remove(tp)
    for a in [k for k in sp.attrib]:
        del sp.attrib[a]
    return sp


def _section_groups(doc):
    """구역마다 블록 묶음(구역은 '문단 안 sectPr' 로 끝나고, 마지막은 본문 sectPr)."""
    groups, cur = [], []
    for el in doc.element.body:
        if el.tag not in (qn("w:p"), qn("w:tbl")):
            continue
        cur.append(el)
        if el.tag == qn("w:p") and el.find(qn("w:pPr") + "/" + qn("w:sectPr")) is not None:
            groups.append(cur)
            cur = []
    groups.append(cur)
    return groups


def _holder(sect):
    """구역 끝 문단(문단 안 sectPr)."""
    from docx.oxml.parser import OxmlElement            # python-docx 요소로 만들어야 문단 기능이 붙는다
    p = OxmlElement("w:p")
    ppr = OxmlElement("w:pPr")
    p.append(ppr)
    ppr.append(sect)
    return p


def _place_bond(doc, tpl, bond_els, place):
    """사모사채 개요를 넣는다(2026-10-08 사용자 확정).

    place = None          → 표지 바로 다음 쪽(하이라이트를 못 찾았을 때)
    place = (k, same)     → 본문 k 번째 구역(=원본 하이라이트 마지막 쪽) 뒤.
                            same=True 면 그 쪽 남은 여백에 이어서, False 면 새 쪽.
    """
    body = doc.element.body
    groups = [g for g in _section_groups(doc) if g]
    if place is None or not groups:
        first = groups[0][0] if groups else body.find(qn("w:sectPr"))
        for el in bond_els:
            first.addprevious(el)
        first.addprevious(_holder(_sect_from_template(tpl, title_page=False)))
        return
    k, same = place
    k = min(k, len(groups) - 1)
    grp = groups[k]
    last = grp[-1]
    ends_section = last.tag == qn("w:p") and last.find(qn("w:pPr") + "/" + qn("w:sectPr")) is not None
    if same:
        # 그 쪽 끝의 빈 문단(pdf2docx 가 쪽 아래를 채운 것)을 걷어내고 바로 뒤에 잇는다
        for el in reversed(grp[:-1] if ends_section else grp):
            if el.tag == qn("w:p") and not "".join(el.itertext()).strip() \
                    and el.find(".//" + qn("w:drawing")) is None:
                body.remove(el)
            else:
                break
        anchor = last if ends_section else None
        for el in bond_els:
            if anchor is not None:
                anchor.addprevious(el)
            else:
                body.find(qn("w:sectPr")).addprevious(el)
        return
    if ends_section:
        for el in reversed(bond_els + [_holder(_sect_from_template(tpl, title_page=False))]):
            last.addnext(el)
    else:
        # 마지막 구역 뒤 — 지금 구역을 닫고(본문 sectPr 복사) 그 뒤에 개요를 둔다
        sp = body.find(qn("w:sectPr"))
        sp.addprevious(_holder(copy.deepcopy(sp)))
        for el in bond_els:
            sp.addprevious(el)


def build(src_bytes, *, title_lines, header_text, date_text, photo, bond, bond_place=None, tpl_path=None):
    """(PDF 에서 만든) 본문 워드 bytes → 회사양식 워드 bytes, 보고(dict).

    본문은 pdf_input.convert 가 이미 회사 규격(A4·여백)으로 맞춰 둔 것이다.
    bond_place : _place_bond 참고.
    """
    tpl_path = tpl_path or template_path()
    tpl = docx.Document(tpl_path)
    doc = docx.Document(io.BytesIO(src_bytes))
    rep = {}

    apply_fonts(doc)
    rep["머리글 칸"] = paint_headers(doc)

    body = doc.element.body
    # 원본 구역들: 쪽번호를 처음부터 다시 세는 설정 지우기(표지부터 이어 세게)
    for sp in body.iter(qn("w:sectPr")):
        pn = sp.find(qn("w:pgNumType"))
        if pn is not None:
            pn.attrib.pop(qn("w:start"), None)

    # ── 사모사채 개요: 하이라이트 뒤 ──
    bond_els = _build_bond_section(doc, bond)
    for el in bond_els:
        body.remove(el)
    if bond_place is not None and bond_place[1]:
        h = bond_els[0]                                     # 같은 쪽에 이을 때는 위와 조금 띄운다
        ppr = _first_child(h, "pPr")
        _put(ppr, ppr.makeelement(qn("w:spacing"), {qn("w:before"): "480", qn("w:after"): "200"}))
    _place_bond(doc, tpl, bond_els, bond_place)
    rep["개요 자리"] = ("표지 다음 쪽" if bond_place is None else
                      f"하이라이트 뒤 {'같은 쪽에 이어서' if bond_place[1] else '새 쪽'}")

    # ── 맨 뒤: 연락처 쪽(세로 A4 구역) ──
    old_last = body.find(qn("w:sectPr"))
    holder = etree.SubElement(body, qn("w:p"))
    body.remove(holder)
    ppr = etree.SubElement(holder, qn("w:pPr"))
    ppr.append(copy.deepcopy(old_last))
    body.insert(list(body).index(old_last), holder)          # 원본 마지막 구역을 문단 구역나눔으로
    body.replace(old_last, _sect_from_template(tpl, title_page=False))
    contacts = read_contacts(tpl)
    for el in _build_contact_section(doc, contacts):
        body.remove(el)
        body.insert(list(body).index(body.find(qn("w:sectPr"))), el)

    # ── 맨 앞: 회사 표지(혼자 한 쪽) ──
    _merge_nsmap(doc, tpl)
    body = doc.element.body
    cover = _cover_elements(tpl, doc, title_lines, date_text, photo)
    # 양식 표지 두 번째 문단 끝의 '쪽 나눔' 은 뺀다 — 구역 끝 문단이 새 쪽을 연다(빈 쪽 방지)
    for br in list(cover[-1].iter(qn("w:br"))):
        if br.get(qn("w:type")) == "page":
            br.getparent().remove(br)
    for i, el in enumerate(cover + [_holder(_sect_from_template(tpl, title_page=True))]):
        body.insert(i, el)

    # ── 머리말·꼬리말: 첫 구역에만 넣고 나머지는 '앞과 같게' ──
    tpl_sec = tpl.sections[0]
    secs = doc.sections
    for k, s in enumerate(secs):
        s.different_first_page_header_footer = (k == 0)
        for hf in (s.header, s.footer, s.even_page_header, s.even_page_footer):
            hf.is_linked_to_previous = True
        if k > 0:
            s.first_page_header.is_linked_to_previous = True
            s.first_page_footer.is_linked_to_previous = True
    s0 = secs[0]
    s0.header.is_linked_to_previous = False
    s0.footer.is_linked_to_previous = False
    _copy_hdrftr(tpl, doc, tpl_sec.header.part, s0.header, header_text)
    _copy_hdrftr(tpl, doc, tpl_sec.footer.part, s0.footer)
    s0.first_page_header.is_linked_to_previous = False      # 표지는 머리말·꼬리말 없음
    s0.first_page_footer.is_linked_to_previous = False
    for hf in (s0.first_page_header, s0.first_page_footer):
        for p in hf.paragraphs[1:]:
            p._p.getparent().remove(p._p)
        if hf.paragraphs:
            for r in list(hf.paragraphs[0].runs):
                r._r.getparent().remove(r._r)
    # 홀짝 머리말 설정 끄기
    st = doc.settings.element
    eo = st.find(qn("w:evenAndOddHeaders"))
    if eo is not None:
        st.remove(eo)

    _renumber_shapes(doc)

    out = io.BytesIO()
    doc.save(out)
    rep["쪽 구역 수"] = len(doc.sections)
    return out.getvalue(), rep


def default_date_text(today=None):
    d = today or date.today()
    return d.strftime("%B %Y")
