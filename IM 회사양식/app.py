# -*- coding: utf-8 -*-
"""IM 회사양식 — 원본 IM(워드·PDF)을 내용 그대로 두고 회사 양식만 입혀 워드로.

바꾸는 것   : 표지 · 위 머리말 · 아래 회사 로고 · 글꼴(Pretendard) · 표 머리글 색(네이비) · 맨 끝 연락처 쪽
더하는 것   : Ⅰ. 사모사채 개요
안 바꾸는 것: 내용 · 글씨 크기 · 쪽 방향(세로는 세로, 가로는 가로)
결과        : 워드(.docx). 원본이 PDF 면 pdf2docx 로 먼저 워드로 바꾼 뒤 같은 처리를 한다.

양식은 사용자가 만든 `양식\\형식_1.docx`(2026-10-07 확정). 엔진은 engine.py · pdf_input.py.

★세 도구와 한 앱에서 돈다(IM\\app.py 가 runpy 로 실행). 화면 상태 이름이 다른 도구와
  겹치면 값이 섞이므로 전부 'cf_' 를 앞에 붙인다(요약본도 'step' 을 쓴다).
★입력칸 값은 'w_cf_…'(입력칸) 와 'cf_…'(보관) 를 따로 둔다. 스트림릿은 화면에서 사라진 입력칸의
  값을 지우므로, 입력칸 이름으로만 두면 2단계에서 적은 값이 3단계에서 버튼을 누르는 순간 날아간다.

혼자 돌릴 때 :  streamlit run app.py --server.port 8620
"""
import io
import os
import re
import zipfile

import streamlit as st

HERE = os.path.dirname(os.path.abspath(__file__))
STEP_NAMES = ["원본 올리기", "표지·사모사채 개요", "만들기"]

# 사모사채 개요 — 구분 칸은 딜과 상관없이 늘 같다(천안·대전·넷마블 회사 제안서 실측).
BOND_ROWS = ["사모사채명", "사채유형", "발행인", "기초자산", "발행금액",
             "발행일", "만기일", "금융조건", "이자지급주기"]
SS = st.session_state

try:      # 통합 앱이 먼저 불렀으면 두 번 부를 수 없다
    st.set_page_config(page_title="IM 회사양식", page_icon="📝", layout="wide")
except Exception:
    pass


# ── 단계 표시줄 (다른 IM 도구와 같은 모양) ──────────
def render_stepper(current):
    st.markdown("""
    <style>
    .stp-wrap{display:flex;align-items:center;justify-content:center;flex-wrap:wrap;
              gap:2px;padding:12px 0 16px 0;}
    .stp-item{display:flex;flex-direction:column;align-items:center;min-width:82px;max-width:110px;}
    .stp-circle{width:32px;height:32px;border-radius:50%;display:flex;align-items:center;
                justify-content:center;font-size:13px;font-weight:bold;margin-bottom:4px;}
    .stp-done{background:#1a7f37;color:#fff;}
    .stp-active{background:#08377C;color:#fff;box-shadow:0 0 0 3px #cdd4e6;}
    .stp-todo{background:#e5e7eb;color:#9ca3af;}
    .stp-label{font-size:11px;text-align:center;line-height:1.3;word-break:keep-all;}
    .stp-l-done{color:#1a7f37;font-weight:600;}
    .stp-l-active{color:#08377C;font-weight:700;}
    .stp-l-todo{color:#9ca3af;}
    .stp-conn{width:26px;height:2px;margin-bottom:20px;}
    .stp-c-done{background:#1a7f37;} .stp-c-todo{background:#e5e7eb;}
    </style>
    """, unsafe_allow_html=True)
    html = '<div class="stp-wrap">'
    for i, name in enumerate(STEP_NAMES):
        n = i + 1
        if n < current:
            cc, lc, tx, conn = "stp-done", "stp-l-done", "✓", "stp-c-done"
        elif n == current:
            cc, lc, tx, conn = "stp-active", "stp-l-active", str(n), "stp-c-todo"
        else:
            cc, lc, tx, conn = "stp-todo", "stp-l-todo", str(n), "stp-c-todo"
        html += (f'<div class="stp-item"><div class="stp-circle {cc}">{tx}</div>'
                 f'<div class="stp-label {lc}">{name}</div></div>')
        if n < len(STEP_NAMES):
            html += f'<div class="stp-conn {conn}"></div>'
    st.markdown(html + "</div>", unsafe_allow_html=True)
    st.markdown("---")


def _goto(n):
    SS["cf_step"] = n
    st.rerun()


def _field(label, store, default="", help=None, area=False):
    """보관값(store)을 바탕으로 입력칸을 그리고, 바뀐 값을 다시 보관한다."""
    SS.setdefault(store, default)
    fn = st.text_area if area else st.text_input
    val = fn(label, value=SS[store], key="w_" + store, help=help)
    SS[store] = val
    return val


def _read_source(name, data):
    """올린 파일의 종류·쪽수·쪽 방향을 잰다. 내용은 건드리지 않는다."""
    ext = os.path.splitext(name)[1].lower()
    info = {"kind": "PDF" if ext == ".pdf" else "워드", "pages": None, "orient": "?"}
    if ext == ".pdf":
        import fitz
        doc = fitz.open(stream=data, filetype="pdf")
        info["pages"] = doc.page_count
        tall = sum(1 for p in doc if p.rect.height >= p.rect.width)
        info["orient"] = ("세로" if tall == doc.page_count else
                          "가로" if tall == 0 else f"섞임(세로 {tall}쪽)")
    else:
        import docx
        d = docx.Document(io.BytesIO(data))
        tall = sum(1 for s in d.sections if s.page_height >= s.page_width)
        n = len(d.sections)
        info["orient"] = ("세로" if tall == n else "가로" if tall == 0 else "섞임")
        info["tables"] = len(d.tables)
    return info


def _prepare(name, data):
    """원본에서 표지 제목·사진 후보·원본 표지/끝쪽 위치를 뽑아 보관한다(올릴 때 한 번)."""
    import engine
    from datetime import date
    info = SS["cf_info"]
    if info["kind"] == "PDF":
        import pdf_input
        info["text_per_page"] = pdf_input.text_per_page(data)
        title = pdf_input.suggest_title(data)
        photos = pdf_input.cover_photo_candidates(data)
        import fitz
        info["last_contact"] = pdf_input.last_page_is_contact(fitz.open(stream=data, filetype="pdf"))
    else:
        import docx
        import pages
        d = docx.Document(io.BytesIO(data))
        end, why, cover_txt = engine.find_cover(d)
        info["cover_end"], info["cover_why"], info["cover_txt"] = end, why, cover_txt
        bl = pages.blocks(d.element.body)
        info["cover_choices"] = [(i, re.sub(r"\s+", " ", pages.text_of(e)).strip()[:40])
                                 for i, e in enumerate(bl[:80]) if pages.text_of(e).strip()]
        start, tail_txt, _ = engine.find_contact_tail(d)
        info["tail_start"], info["tail_txt"] = start, re.sub(r"\s+", " ", tail_txt).strip()[:200]
        title = engine.suggest_title(d, end)
        photos = engine.photo_candidates(d)
    title = (title + ["", ""])[:2]
    SS["cf_title1"], SS["cf_title2"] = title
    base = re.sub(r"^\[?IM\]?[\s\-_]*", "", os.path.splitext(name)[0]).strip() or title[0]
    SS["cf_header"] = f"{base} 사모사채 제안서 {date.today().year}"
    SS["cf_date"] = engine.default_date_text()
    SS["cf_photos"] = photos
    SS["cf_photo_pick"] = 0 if photos else -1
    SS.pop("cf_photo_custom", None)
    for k in ("cf_out", "cf_rep", "cf_build_err"):
        SS.pop(k, None)


SS.setdefault("cf_step", 1)
step = SS["cf_step"]

st.markdown("""
<style>
.block-container{padding-top:1.6rem;max-width:100%;
                 padding-left:2.2rem;padding-right:2.2rem;}
h1,h2,h3{color:#08377C;}
div.stButton>button{border-radius:4px;border:1px solid #08377C;color:#08377C;
                    background:#fff;font-weight:700;}
div.stButton>button:hover{background:#eef3fa;color:#08377C;border-color:#08377C;}
div.stButton>button[kind="primary"]{background:#1A2B5E;color:#fff;border-color:#1A2B5E;}
div.stButton>button[kind="primary"]:hover{background:#08377C;border-color:#08377C;}
div.stDownloadButton>button{background:#08377C;color:#fff;border:1px solid #08377C;
                            border-radius:4px;font-weight:700;}
div.stDownloadButton>button:hover{background:#0063A1;border-color:#0063A1;color:#fff;}
hr{border-color:#dfe6f0;}
</style>
""", unsafe_allow_html=True)

st.markdown("<h1 style='margin:0 0 .1rem 0;line-height:1.1;'>IM 회사양식</h1>",
            unsafe_allow_html=True)
st.caption("받은 IM 을 **내용 그대로** 두고, 회사 양식(표지·머리말·로고·글꼴·표 머리글 색·연락처)만 "
           "입혀 **워드**로 만듭니다. 사모사채 개요도 넣습니다.")
render_stepper(step)


# ══════════════════════════════════════════════════
# 1단계 — 원본 올리기
# ══════════════════════════════════════════════════
if step == 1:
    st.markdown("### 1단계 · 원본 IM 올리기")
    st.caption("받은 IM 을 그대로 올리세요. **워드(.docx)와 PDF 둘 다** 됩니다. "
               "세로 문서는 세로 그대로, 가로 문서는 가로 그대로 나옵니다.")

    c1, c2 = st.columns(2)
    c1.success("**그대로 두는 것**\n\n- 본문 내용·표·그림\n- 글씨 크기\n- 쪽 방향(세로/가로)")
    c2.info("**회사 양식으로 바꾸는 것**\n\n- 표지\n- 위쪽 머리말 · 아래쪽 로고(증권사 것을 빼고 회사 것으로)\n"
            "- 글꼴(Pretendard)\n- 표 머리글 색(네이비)\n- 맨 끝 연락처 쪽\n\n**더하는 것** : Ⅰ. 사모사채 개요")

    with st.expander("📎 워드와 PDF 중 무엇을 올릴까요?"):
        st.markdown(
            "- **워드 원본이 있으면 워드를 올려 주세요.** 가장 깔끔하게 나옵니다.\n"
            "- **PDF 밖에 없으면 PDF 를 올리면 됩니다.** 도구가 먼저 워드로 바꾼 뒤 양식을 입힙니다. "
            "이때 표나 줄바꿈이 원본과 조금 다를 수 있어, 받은 뒤 한 번 훑어봐 주세요. "
            "PDF 는 30쪽 기준 1~2분쯤 걸립니다.\n"
            "- 글자를 마우스로 드래그해도 선택되지 않는 PDF(스캔본·글자를 그림으로 바꾼 것)는 바꿀 수 없습니다.\n"
            "- 옛날 워드 형식(.doc)은 워드에서 열어 **F12 → 파일 형식 'Word 문서(*.docx)'** 로 "
            "저장한 뒤 올려 주세요.")

    up = st.file_uploader("원본 IM (워드 또는 PDF)", type=["docx", "pdf"], key="w_cf_up")
    if up is not None and SS.get("cf_name") != up.name:
        SS["cf_name"] = up.name
        SS["cf_bytes"] = up.getvalue()
        for k in ("cf_err", "cf_repair"):
            SS.pop(k, None)
        try:
            SS["cf_info"] = _read_source(up.name, up.getvalue())
        except zipfile.BadZipFile:
            # ★실제로 있었다 — 부산 남포동 워드 원본은 파일 끝이 잘려 있어 못 열었다.
            #   먼저 스스로 고쳐 보고(docx_repair), 그래도 안 될 때만 사람에게 부탁한다.
            from docx_repair import repair_docx, missing_summary
            fixed, missing = repair_docx(up.getvalue())
            info = None
            if fixed is not None:
                try:
                    info = _read_source(up.name, fixed)
                except Exception:
                    info = None
            if info is not None:
                SS["cf_bytes"] = fixed          # 이후 단계는 고친 파일로 진행
                SS["cf_info"] = info
                SS["cf_repair"] = missing_summary(missing)
            else:
                SS["cf_info"] = None
                SS["cf_err"] = "broken"
        except Exception as e:
            SS["cf_info"] = None
            SS["cf_err"] = f"파일을 열지 못했습니다 : {e}"
        if SS.get("cf_info"):
            with st.spinner("원본을 살펴보는 중…"):
                try:
                    _prepare(up.name, SS["cf_bytes"])
                except Exception as e:
                    SS["cf_info"] = None
                    SS["cf_err"] = f"원본을 읽다가 멈췄습니다 : {e}"

    err = SS.get("cf_err")
    if err == "broken":
        st.error(
            "**이 워드 파일은 깨져 있어서 열 수 없습니다.** 자동으로 고쳐 보았지만 본문까지 "
            "잘려 나가 되살리지 못했습니다.\n\n"
            "**왜 이런가요?** 워드 파일(.docx)은 본문·서식·그림 같은 여러 조각을 하나로 묶은 "
            "압축 파일이고, 맨 끝에 '어느 조각이 어디 있는지' 적힌 **목차**가 붙어 있습니다. "
            "메일로 주고받거나 내려받는 도중에 **파일 끝부분이 잘리면** 이 목차가 사라져 "
            "프로그램이 파일을 열지 못합니다. 워드 프로그램은 이런 파일도 스스로 고쳐서 열어 주기 때문에, "
            "내 컴퓨터에서는 멀쩡해 보일 수 있습니다.\n\n"
            "**이렇게 해 주세요**\n"
            "1. 이 파일을 **워드에서 열어** 주세요. '복구할까요?' 라고 물으면 **예**를 누릅니다.\n"
            "2. **F12** → 파일 형식 **'Word 문서(*.docx)'** → 저장.\n"
            "3. 새로 저장한 파일을 여기에 다시 올려 주세요.\n\n"
            "같은 IM 의 **PDF** 가 있으면 PDF 를 올리는 것이 더 정확합니다. 보낸 쪽에 원본을 "
            "다시 받을 수 있으면 그게 가장 좋습니다.")
    elif err:
        st.error(err)

    repair = SS.get("cf_repair")
    if repair is not None:
        if repair:
            st.warning(
                "🛠️ **깨진 워드 파일이라 자동으로 고쳐서 열었습니다.** 다만 **되살리지 못한 부분**이 있습니다 : "
                f"**{repair}**\n\n"
                "**왜 이런가요?** 이 파일은 받는 도중에 **파일 끝부분이 잘려 나간** 상태입니다. 워드 파일은 "
                "본문·서식·그림 같은 조각을 묶어 둔 것인데, 잘려 나간 뒤쪽에 있던 조각은 파일 안에 **아예 "
                "남아 있지 않아** 어떤 프로그램으로도 되살릴 수 없습니다(워드에서 다시 저장해도 마찬가지입니다).\n\n"
                "**그래서 결과물이 원본과 이렇게 다를 수 있습니다**\n"
                "- 서식이 빠지면 글머리표(✓ 등)가 번호로 바뀌거나 줄 간격이 넓어져 쪽이 밀립니다.\n"
                "- 빠진 그림은 그 자리가 비어 나옵니다.\n\n"
                "**더 정확하게 하려면** 같은 IM 의 **PDF** 를 올리거나, 보낸 쪽에 **원본 워드를 다시 요청**해 주세요. "
                "이대로 계속 진행해도 됩니다.")
        else:
            st.info("🛠️ 깨진 워드 파일이라 자동으로 고쳐서 열었습니다. 빠진 내용은 없습니다.")

    info = SS.get("cf_info")
    scanned = bool(info and info["kind"] == "PDF" and info.get("text_per_page", 999) < 100)
    if info:
        st.success(f"올린 파일 — **{SS['cf_name']}**")
        m1, m2, m3 = st.columns(3)
        m1.metric("종류", info["kind"])
        m2.metric("쪽 방향", info["orient"])
        if info["pages"] is not None:
            m3.metric("쪽수", f"{info['pages']}쪽")
        else:
            m3.metric("표 개수", f"{info.get('tables', 0)}개")
        if scanned:
            st.error(
                "🚫 **이 PDF 는 글자를 읽을 수 없습니다** (쪽당 글자 "
                f"{info['text_per_page']:.0f}자 — 보통 IM 은 600~1,100자).\n\n"
                "스캔본이거나 글자를 그림으로 바꿔 저장한 PDF 입니다. 겉으로는 글자가 또렷해도 컴퓨터는 "
                "그림으로만 봅니다(실제 사례: 신사동 평화빌딩 PDF). 이런 PDF 는 워드로 바꿀 수 없습니다.\n\n"
                "**워드 원본**이나 **글자가 살아 있는 PDF**(워드에서 F12 로 PDF 저장한 것)를 받아 올려 주세요.")
        elif info["kind"] == "PDF":
            st.caption("PDF 는 3단계에서 먼저 워드로 바꾼 뒤 양식을 입힙니다.")

        if not scanned:
            st.markdown("---")
            _b1, _b2, _b3 = st.columns([2, 4, 2])
            if _b3.button("다음 단계 →", type="primary", use_container_width=True):
                _goto(2)


# ══════════════════════════════════════════════════
# 2단계 — 표지 · 사모사채 개요
# ══════════════════════════════════════════════════
elif step == 2:
    st.markdown("### 2단계 · 표지 · 사모사채 개요")
    if not SS.get("cf_info"):
        st.warning("먼저 1단계에서 원본을 올려 주세요.")
    else:
        st.markdown("#### 표지 · 머리말")
        st.caption("원본 표지에서 제목을 찾아 미리 채워 두었습니다. 고칠 곳만 고치세요.")
        a, b = st.columns(2)
        with a:
            _field("표지 제목 (첫째 줄)", "cf_title1")
            _field("표지 제목 (둘째 줄 · 없으면 비워 두세요)", "cf_title2")
        with b:
            _field("위쪽 머리말 문구", "cf_header",
                   help="본문 모든 쪽 맨 위에 들어가는 문구입니다. 예) 신촌지역 브릿지 대출 사모사채 제안서 2026")
            _field("표지 날짜", "cf_date", help="영어로 적습니다. 예) October 2026")

        st.markdown("#### 표지 사진")
        photos = SS.get("cf_photos") or []
        if photos:
            st.caption("원본에 있는 사진 중 표지에 쓸 만한 것을 골라 두었습니다(색이 다양한 큰 실사 사진 먼저). "
                       "표지 사진 칸 비율(가로 2 : 세로 1)에 맞춰 가운데를 잘라 넣습니다.")
            cols = st.columns(min(len(photos), 6))
            for i, (col, blob) in enumerate(zip(cols, photos)):
                with col:
                    st.image(blob, use_container_width=True)
                    picked = SS.get("cf_photo_pick") == i
                    if st.button("✅ 선택됨" if picked else "이 사진 쓰기", key=f"w_cf_pick_{i}",
                                 type="primary" if picked else "secondary", use_container_width=True):
                        SS["cf_photo_pick"] = i
                        st.rerun()
        else:
            st.caption("원본에서 표지에 쓸 사진을 찾지 못했습니다.")
        o1, o2 = st.columns([1, 2])
        with o1:
            none_picked = SS.get("cf_photo_pick") == -1
            if st.button("✅ 사진 없이" if none_picked else "사진 없이 만들기", key="w_cf_pick_none",
                         type="primary" if none_picked else "secondary", use_container_width=True):
                SS["cf_photo_pick"] = -1
                st.rerun()
        with o2:
            mine = st.file_uploader("다른 사진 직접 올리기 (jpg·png)", type=["jpg", "jpeg", "png"],
                                    key="w_cf_photo_up")
            if mine is not None:
                SS["cf_photo_custom"] = mine.getvalue()
                SS["cf_photo_pick"] = -2
            if SS.get("cf_photo_pick") == -2 and SS.get("cf_photo_custom"):
                st.image(SS["cf_photo_custom"], caption="✅ 직접 올린 사진을 씁니다", width=260)

        st.markdown("---")
        st.markdown("#### Ⅰ. 사모사채 개요")
        st.caption("왼쪽 구분은 항상 같고, 오른쪽 내용만 이번 딜에 맞게 적어 주세요. "
                   "비워 둔 칸은 빈칸으로 들어갑니다. (원본에서 자동으로 채우는 기능은 나중에 붙일 예정입니다.)")
        for row in BOND_ROWS:
            k1, k2 = st.columns([1, 4])
            k1.markdown(f"<div style='padding-top:8px;font-weight:700;color:#08377C;'>{row}</div>",
                        unsafe_allow_html=True)
            with k2:
                SS.setdefault(f"cf_bond_{row}", "")
                SS[f"cf_bond_{row}"] = st.text_input(row, value=SS[f"cf_bond_{row}"],
                                                     key=f"w_cf_bond_{row}", label_visibility="collapsed")

    st.markdown("---")
    p1, _sp, p3 = st.columns([2, 4, 2])
    if p1.button("← 이전 단계", use_container_width=True):
        _goto(1)
    if SS.get("cf_info") and p3.button("다음 단계 →", type="primary", use_container_width=True):
        _goto(3)


# ══════════════════════════════════════════════════
# 3단계 — 만들기
# ══════════════════════════════════════════════════
elif step == 3:
    st.markdown("### 3단계 · 만들기")
    info = SS.get("cf_info")
    if not SS.get("cf_bytes") or not info:
        st.warning("먼저 1단계에서 원본을 올려 주세요.")
    else:
        import engine
        is_pdf = info["kind"] == "PDF"
        st.write(f"**{SS.get('cf_name', '')}** · {info.get('kind', '')} · {info.get('orient', '')}")
        filled = [r for r in BOND_ROWS if (SS.get(f"cf_bond_{r}") or "").strip()]
        st.write(f"표지 제목 : **{SS.get('cf_title1', '')} {SS.get('cf_title2', '')}** · "
                 f"사모사채 개요 : {len(filled)}/{len(BOND_ROWS)}칸 채움")

        st.markdown("#### 원본에서 뺄 쪽")
        if is_pdf:
            SS.setdefault("cf_drop_cover", True)
            SS["cf_drop_cover"] = st.checkbox("원본 1쪽(증권사 표지)을 빼고 회사 표지로 바꾸기",
                                              value=SS["cf_drop_cover"], key="w_cf_drop_cover")
            SS.setdefault("cf_drop_tail", bool(info.get("last_contact")))
            SS["cf_drop_tail"] = st.checkbox(
                "원본 마지막 쪽(증권사 연락처)을 빼기"
                + (" — 메일·전화번호가 있어 연락처 쪽으로 보입니다" if info.get("last_contact")
                   else " — 연락처 쪽으로 보이지 않습니다(본문일 수 있음)"),
                value=SS["cf_drop_tail"], key="w_cf_drop_tail")
        else:
            SS.setdefault("cf_drop_cover", info.get("cover_end") is not None)
            SS["cf_drop_cover"] = st.checkbox("원본 표지를 빼고 회사 표지로 바꾸기", value=SS["cf_drop_cover"],
                                              key="w_cf_drop_cover")
            choices = info.get("cover_choices") or []
            if SS["cf_drop_cover"] and choices:
                idxs = [i for i, _ in choices]
                auto = info.get("cover_end")
                default = max([i for i in idxs if auto is not None and i <= auto] or [idxs[0]])
                SS.setdefault("cf_cover_end", default)
                labels = {i: f"{t}" for i, t in choices}
                st.caption("워드 파일에는 '쪽' 이 적혀 있지 않아 원본 표지가 어디서 끝나는지 어림했습니다. "
                           "아래 글이 **원본 표지의 마지막 글**이 맞는지 확인해 주세요. 틀리면 바꿔 주세요.")
                pos = idxs.index(SS["cf_cover_end"]) if SS["cf_cover_end"] in idxs else 0
                SS["cf_cover_end"] = st.selectbox("원본 표지는 여기까지", idxs, index=pos,
                                                  format_func=lambda i: labels.get(i, str(i)),
                                                  key="w_cf_cover_end")
            if info.get("tail_start") is not None:
                SS.setdefault("cf_drop_tail", True)
                SS["cf_drop_tail"] = st.checkbox("원본 마지막 쪽(증권사 연락처)을 빼기", value=SS["cf_drop_tail"],
                                                 key="w_cf_drop_tail")
                st.caption(f"마지막 쪽 내용 : {info.get('tail_txt', '')[:120]}…")
            else:
                SS["cf_drop_tail"] = False

        st.markdown("#### 파일 이름")
        _base = os.path.splitext(SS.get("cf_name") or "IM")[0]
        SS.setdefault("cf_fname", f"{_base}_회사양식")
        f1, f2 = st.columns([5, 1])
        with f1:
            _field("파일명", "cf_fname", help="내려받을 파일 이름입니다. 확장자(.docx)는 자동으로 붙습니다.")
        f2.markdown("<div style='padding-top:34px;color:#5b6b85;'>.docx</div>", unsafe_allow_html=True)

        tpl = engine.template_path()
        if not tpl:
            st.error("회사 양식 파일(양식\\형식_1.docx)을 찾지 못해 만들 수 없습니다. 관리자에게 알려 주세요.")
        elif st.button("📝 회사 양식 워드 만들기", type="primary"):
            for k in ("cf_out", "cf_rep", "cf_build_err"):
                SS.pop(k, None)
            pick = SS.get("cf_photo_pick", -1)
            photo = (SS.get("cf_photo_custom") if pick == -2 else
                     (SS.get("cf_photos") or [None])[pick] if pick is not None and pick >= 0 else None)
            bond = {r: SS.get(f"cf_bond_{r}", "") for r in BOND_ROWS}
            title_lines = [SS.get("cf_title1", ""), SS.get("cf_title2", "")]
            try:
                with st.spinner("만드는 중입니다… (PDF 는 30쪽 기준 1~2분)"):
                    src = SS["cf_bytes"]
                    rep_pdf = {}
                    if is_pdf:
                        import pdf_input
                        src, rep_pdf = pdf_input.pdf_to_docx(src, drop_cover=SS.get("cf_drop_cover", True),
                                                             drop_last=SS.get("cf_drop_tail", False))
                        cover_end, drop_tail = None, False
                    else:
                        cover_end = SS.get("cf_cover_end") if SS.get("cf_drop_cover") else None
                        drop_tail = SS.get("cf_drop_tail", False)
                    out, rep = engine.build(src, title_lines=title_lines, header_text=SS.get("cf_header", ""),
                                            date_text=SS.get("cf_date", ""), photo=photo, bond=bond,
                                            drop_tail=drop_tail, cover_end=cover_end)
                SS["cf_out"] = out
                SS["cf_rep"] = {**rep, **{k: v for k, v in rep_pdf.items() if v}}   # PDF 쪽에서 뺀 것도 보고에 남게
            except Exception:
                import traceback
                SS["cf_build_err"] = traceback.format_exc()

        if SS.get("cf_build_err"):
            st.error("만들지 못했습니다. 아래 내용을 관리자에게 보내 주세요.")
            st.code(SS["cf_build_err"])

        if SS.get("cf_out"):
            rep = SS.get("cf_rep") or {}
            st.success("✅ 다 만들었습니다. 아래 버튼으로 내려받으세요.")
            nm = re.sub(r'[\\/:*?"<>|]', "_", (SS.get("cf_fname") or "").strip())[:80] or f"{_base}_회사양식"
            if nm.lower().endswith(".docx"):
                nm = nm[:-5]
            st.download_button("⬇️ 워드 내려받기", data=SS["cf_out"], file_name=f"{nm}.docx",
                               mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document")
            done = [f"표 머리글 칸 {rep.get('머리글 칸', 0)}개를 네이비로 바꿨습니다."]
            if rep.get("원본 표지 뺌"):
                done.append("원본 표지를 빼고 회사 표지를 넣었습니다.")
            if rep.get("원본 끝쪽 뺌"):
                done.append("원본 마지막 쪽(증권사 연락처)을 빼고 회사 연락처 쪽을 넣었습니다.")
            if rep.get("머리말·꼬리말 지운 쪽"):
                done.append(f"PDF 쪽마다 되풀이되던 증권사 머리말·꼬리말을 {rep['머리말·꼬리말 지운 쪽']}쪽에서 지웠습니다.")
            st.markdown("\n".join("- " + t for t in done))
            with st.expander("⚠️ 받은 뒤 확인해 주세요"):
                st.markdown(
                    "- **글꼴 Pretendard** 가 깔린 컴퓨터에서 열어야 같은 모양으로 보입니다. 밖으로 보낼 때는 "
                    "워드에서 **PDF 로 저장해서** 보내면 어디서나 똑같이 보입니다.\n"
                    "- 글꼴이 바뀌어 글자 폭·줄 높이가 조금 달라지므로 **쪽이 넘어가는 자리가 원본과 다를 수 있습니다.** "
                    "표가 두 쪽에 걸치지 않았는지 한 번 훑어봐 주세요.\n"
                    "- 원본의 **기울임(이탤릭) 글씨**는 Pretendard 에 기울임 글꼴이 없어 PDF 로 저장하면 다른 글꼴로 보일 수 있습니다.\n"
                    "- 원본 그림 **안에** 박힌 증권사 로고(구조도 그림 등)는 그림이라 바꿀 수 없습니다.\n"
                    "- PDF 원본은 워드로 바꾸는 과정에서 표·줄바꿈이 조금 다를 수 있습니다.")

    st.markdown("---")
    p1, _sp, _p3 = st.columns([2, 4, 2])
    if p1.button("← 이전 단계", use_container_width=True):
        _goto(2)
